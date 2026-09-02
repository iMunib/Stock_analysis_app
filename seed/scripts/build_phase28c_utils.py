# -*- coding: utf-8 -*-
"""Phase 2.8 A3 utilities detail - H/FTS cash metric value, capex, allowed ROE."""
import openpyxl

XLSX = r"C:\Users\RehmanPC\.openclaw-autoclaw\workspace\na_financials_research\NA_Company_Financials.xlsx"
wb = openpyxl.load_workbook(XLSX)
ws = wb["Utility_Pipeline"]
hdr = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(hdr)}
ticker_rows = {ws.cell(row=r, column=col["Ticker"]).value: r for r in range(2, ws.max_row + 1)}

def setv(tk, field, val):
    if tk in ticker_rows and field in col:
        ws.cell(row=ticker_rows[tk], column=col[field], value=val)

def append_note(tk, text):
    if tk not in ticker_rows:
        return
    r = ticker_rows[tk]
    cur = ws.cell(row=r, column=col["Notes"]).value or ""
    ws.cell(row=r, column=col["Notes"], value=(cur + " " + text).strip())

# Hydro One (H): issuer cash metric value + capex
setv("H", "FFO_or_Equivalent", -671000000)  # net cash from ops 2,695 - capital investments 3,366
setv("H", "Capex", 3366000000)
setv("H", "Source_URL", "https://hydroone.mediaroom.com/2026-02-13-Hydro-One-Reports-Fourth-Quarter-Results")
setv("H", "Confidence", "High (issuer press release)")
append_note("H", "FY2025 net cash from operating activities $2,695M; capital investments $3,366M -> 'cash from operations after capital expenditures' = -$671M (transmission growth capex, negative is growth not distress).")

# Fortis (FTS): allowed ROE
setv("FTS", "Allowed_ROE", 0.095)
append_note("FTS", "Allowed ROE 9.5% (UNS Energy 3-yr rate plan, eff. Jul 2025); jurisdiction-specific (e.g., FortisAlberta 9.28% 2024). Rate base $42.4B midyear 2025, 7% growth.")

wb.save(XLSX)
print("Saved A3 utilities detail (H/FTS).")
