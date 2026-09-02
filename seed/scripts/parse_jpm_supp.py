import openpyxl
p = r"C:\Users\RehmanPC\.openclaw-autoclaw\workspace\na_financials_research\raw\docs\jpm_4q25_supp.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
print("sheets:", wb.sheetnames)
# scan all sheets for key metric labels
keys = ["CET1", "efficiency", "net interest margin", "NIM", "net charge-off", "NCO", "nonperforming", "NPL", "tangible book value", "TBVPS", "common equity Tier 1", "ROTCE"]
for sn in wb.sheetnames:
    ws = wb[sn]
    hits = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(k.lower() in c.value.lower() for k in keys):
                hits.append((c.coordinate, c.value[:70]))
    if hits:
        print(f"\n=== {sn} ===")
        for coord, v in hits[:30]:
            print(f"  {coord}: {v}")
