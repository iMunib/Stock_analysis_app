import os
import sys
import traceback
import zipfile
from datetime import date

import numpy as np
import pandas as pd

ROOT = r"C:\Users\RehmanPC\Documents\NA_Financials_ClaudeCode_Handoff\na_financials_research"
CANON = os.path.join(ROOT, "NA_Company_Financials.xlsx")
OUT = os.path.join(ROOT, "Sector_Financials_Phase1.xlsx")
LOGD = os.path.join(ROOT, "logs")
TODAY = date.today().isoformat()

ID_COLS_UNI = ["Company_ID", "Company_Name", "Primary_Ticker", "Country_of_Listing",
               "Exchange", "GICS_Sector", "GICS_Industry", "Custom_Industry_Sheet",
               "In_SP500", "In_TSX_Composite"]
CORE_KEEP_ID = ["Primary_Currency", "Extraction_Status", "Source_Primary", "Fiscal_Year_End"]
STMT_COLS = ["Revenue", "Net_Income", "Diluted_EPS", "Gross_Profit", "Operating_Cash_Flow",
             "Capex", "Free_Cash_Flow", "Total_Debt", "Book_Equity", "Cash_ST_Investments",
             "Total_Assets", "Total_Liabilities"]
DERIVED_COLS = ["FCF_Reported", "FCF_Calc", "NetDebt_Calc", "FCFMargin_Calc",
                "GrossMargin_Calc", "ROE_Calc", "Calc_Method_Notes"]
OUT_COLS = ID_COLS_UNI + ["Currency"] + CORE_KEEP_ID[:3] + ["Fiscal_Year_End"] + STMT_COLS + DERIVED_COLS
COLLISION_TICKERS = ["ARE", "EFX", "KEY", "L", "PPL", "T", "TKO"]
ROE_LABEL = "ROE = NI / ending equity (not average equity)"


def fail(msg):
    print("BUILD FAILED:", msg)
    if os.path.exists(OUT):
        os.remove(OUT)
        print("Removed partial output:", OUT)
    sys.exit(1)


def num(s):
    return pd.to_numeric(s, errors="coerce")


def build():
    uni = pd.read_excel(CANON, sheet_name="Universe", keep_default_na=False)
    core = pd.read_excel(CANON, sheet_name="Core_Financials", keep_default_na=False)
    dq_old = pd.read_excel(CANON, sheet_name="Data_Quality", keep_default_na=False)

    assert uni["Company_ID"].nunique() == 720 and len(uni) == 720, "universe broken"
    uid_set = set(uni["Company_ID"].astype(str))

    idcols = [c for c in core.columns if c.startswith("Company_ID")]
    print("=== TASK 1A: Core identity columns ===")
    for c in idcols:
        s = core[c].astype(str).str.strip()
        print(f"  {c!r}: nonblank={int((s != '').sum())}/720 unique={s.nunique()} "
              f"exact-match-Universe={int(s.isin(uid_set).sum())}")

    key = core["Company_ID.1"].astype(str).str.strip()
    blank_mask = key == ""
    repaired_ids = []
    if blank_mask.any():
        fallback = core.loc[blank_mask, "Company_ID"].astype(str).str.strip()
        ok = fallback.isin(uid_set)
        assert ok.all(), "backfill candidates not found in Universe"
        key.loc[blank_mask] = fallback
        repaired_ids = sorted(fallback.unique())
    core["_Join_Key"] = key

    assert core["_Join_Key"].ne("").all(), "join key still has blanks"
    assert core["_Join_Key"].nunique() == 720, "join key not unique"
    assert set(core["_Join_Key"]) == uid_set, "core/universe key sets differ"

    dup_names = [c for c in idcols if c != "Company_ID.1"]
    core_keep = core[["_Join_Key"] + CORE_KEEP_ID + STMT_COLS].copy()
    rename_map = {c: f"{c}_Core" for c in CORE_KEEP_ID + STMT_COLS if c in uni.columns}
    core_keep = core_keep.rename(columns=rename_map)

    def col(name):
        return rename_map.get(name, name)
    for c in STMT_COLS:
        core_keep[c] = num(core_keep[c])

    m = uni.merge(core_keep, how="left", left_on="Company_ID", right_on="_Join_Key",
                  validate="one_to_one")
    assert len(m) == 720, f"merge produced {len(m)} rows"
    assert m["_Join_Key"].notna().all(), "ghost/unmatched universe rows after join"
    m = m.drop(columns=["_Join_Key"])

    print("\n=== Identity repair summary ===")
    print("join key: Universe.Company_ID == Core_Financials['Company_ID.1'] (physical col B)")
    print("rows backfilled:", len(repaired_ids), repaired_ids)
    print("dropped legacy Core ID columns:", idcols)

    cur_u = uni.set_index("Company_ID")["Financials_Native_Currency"].astype(str).str.strip()
    cur_c = m.set_index("Company_ID")[col("Primary_Currency")].astype(str).str.strip()
    mism = int((cur_u.reindex(m["Company_ID"]).values != cur_c.values).sum())
    print("currency mismatches Universe vs Core:", mism)
    print("core columns renamed for merge:", rename_map)

    out = pd.DataFrame(index=m.index)
    for c in ID_COLS_UNI:
        out[c] = m[c].astype(str).str.strip()
    out["Currency"] = m[col("Primary_Currency")].astype(str).str.strip()
    for c in ["Extraction_Status", "Source_Primary", "Fiscal_Year_End"]:
        out[c] = m[col(c)].astype(str).str.strip()
    for c in STMT_COLS:
        out[c] = num(m[col(c)])

    banks = out["Custom_Industry_Sheet"].eq("Banks")
    rev, ocf, capex = num(out["Revenue"]), num(out["Operating_Cash_Flow"]), num(out["Capex"])
    debt, cash = num(out["Total_Debt"]), num(out["Cash_ST_Investments"])
    gp, ni, eq = num(out["Gross_Profit"]), num(out["Net_Income"]), num(out["Book_Equity"])
    fcf_rep = num(m[col("Free_Cash_Flow")])

    out["FCF_Reported"] = fcf_rep
    fcf_calc = ocf - capex.abs()
    fcf_calc[(ocf.isna() | capex.isna() | banks)] = np.nan
    out["FCF_Calc"] = fcf_calc
    nd = debt - cash
    nd[debt.isna() | cash.isna()] = np.nan
    out["NetDebt_Calc"] = nd
    fm = fcf_calc / rev.where(rev != 0)
    fm[banks] = np.nan
    out["FCFMargin_Calc"] = fm.round(6)
    gm = gp / rev.where(rev != 0)
    out["GrossMargin_Calc"] = gm.round(6)
    roe = ni / eq.where(eq > 0)
    out["ROE_Calc"] = roe.round(6)

    tol = fcf_rep.abs() * 1e-6 + 1000
    diverges = capex.lt(0) & fcf_calc.notna() & fcf_rep.notna() & ((fcf_calc - fcf_rep).abs() > tol)
    n_div = int(diverges.sum())
    n_neg_capex = int(capex.lt(0).sum())
    n_pos_capex = int(capex.gt(0).sum())
    roe_blank_eq = int((eq.notna() & (eq <= 0)).sum())

    notes = []
    for i in out.index:
        p = []
        if pd.notna(fcf_calc[i]):
            p.append("FCF_Calc = OCF - abs(Capex)")
            if bool(diverges[i]):
                p.append("differs from reported FCF (stored value used signed Capex)")
        if pd.notna(nd[i]):
            p.append("NetDebt_Calc = Total_Debt - Cash" +
                     (" (gross debt minus cash; not a bank capital metric)" if banks[i] else ""))
        if pd.notna(fm[i]):
            p.append("FCFMargin_Calc = FCF_Calc / Revenue")
        if pd.notna(gm[i]):
            p.append("GrossMargin_Calc = Gross_Profit / Revenue")
        if pd.notna(roe[i]):
            p.append(ROE_LABEL)
        if banks[i]:
            p.append("Banks: industrial cash-flow metrics intentionally blank")
        notes.append("; ".join(p))
    out["Calc_Method_Notes"] = notes

    empty_price = {c: int(num(core[c]).notna().sum()) for c in ["Price", "Market_Cap", "Shares_Diluted"]}
    assert all(v == 0 for v in empty_price.values()), f"unexpected market data present: {empty_price}"

    assert len(out) == 720
    assert out["Company_ID"].nunique() == 720
    assert out["Company_ID"].eq("").sum() == 0

    coll_ids = [f"US:{t}:US" for t in COLLISION_TICKERS] + [f"CA:{t}:TSX" for t in COLLISION_TICKERS]
    coll = out[out["Company_ID"].isin(coll_ids)].sort_values(["Primary_Ticker", "Country_of_Listing"])
    assert len(coll) == 14, f"expected 14 collision rows, got {len(coll)}"

    cov_fields = [("Revenue", "native currency"), ("Net_Income", "native currency"),
                  ("Diluted_EPS", ""), ("Operating_Cash_Flow", "OCF"), ("Capex", ""),
                  ("FCF_Reported", "copy of Core Free_Cash_Flow"), ("FCF_Calc", ROE_LABEL.replace(ROE_LABEL, "OCF - abs(Capex); blank for Banks")),
                  ("Total_Debt", ""), ("Equity (Book_Equity)", "ending balance"), ("Cash (Cash_ST_Investments)", ""),
                  ("NetDebt_Calc", "Total_Debt - Cash; banks included, labeled gross-only"),
                  ("FCFMargin_Calc", "blank for Banks"), ("GrossMargin_Calc", "Gross_Profit / Revenue"),
                  ("ROE_Calc", ROE_LABEL)]
    cov_rows = []
    for label, note in cov_fields:
        src = {"Equity (Book_Equity)": "Book_Equity", "Cash (Cash_ST_Investments)": "Cash_ST_Investments"}.get(label, label)
        n = int(num(out[src]).notna().sum())
        cov_rows.append({"Field": label, "n_non_null": n, "pct": round(n / 7.2, 1), "notes": note})
    cov = pd.DataFrame(cov_rows)

    new_dq = [
        {"Company_ID": "CA:EFX:TSX", "Ticker": "EFX.TO", "Field": "Company_ID (join key)",
         "Issue": "BLANK_JOIN_KEY_REPAIRED", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Repaired in Phase 1 output only",
         "Notes": "Core Company_ID.1 (physical col B) was blank for Enerflex; its four sibling Core ID columns all read CA:EFX:TSX matching Universe; key backfilled from Universe.Company_ID; row retained with real financials (Revenue 2571000000 CAD)."},
        {"Company_ID": "(workbook)", "Ticker": "", "Field": "Core identity columns",
         "Issue": "LEGACY_ID_COLUMNS_DROPPED", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Excluded from output",
         "Notes": "Core columns A/.2/.3/.4 use exchange-suffix IDs (e.g. US:XOM:NYSE) that disagree with the permanent Universe scheme; Universe.Company_ID is the sole key."},
        {"Company_ID": "(workbook)", "Ticker": "", "Field": "Capex",
         "Issue": "SIGN_CONVENTION_MIXED", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Documented; FCF_Calc uses abs(Capex)",
         "Notes": f"{n_pos_capex} rows positive-spend storage, {n_neg_capex} rows negative-outflow storage. Stored Free_Cash_Flow equals OCF minus signed Capex on all {int(num(core['Free_Cash_Flow']).notna().sum())} populated rows, inflating FCF above OCF on the {n_div} negative-capex rows (e.g. ENB stored FCF 21435000000 vs OCF 12270000000)."},
        {"Company_ID": "(workbook)", "Ticker": "", "Field": "FCF_Calc / FCFMargin_Calc",
         "Issue": "BANK_EXCLUSION_BY_DESIGN", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Left blank for Banks",
         "Notes": f"{int(banks.sum())} Banks rows excluded from industrial cash-flow metrics."},
        {"Company_ID": "(workbook)", "Ticker": "", "Field": "Price / Market_Cap / Shares_Diluted",
         "Issue": "COLUMNS_OMITTED_NO_INPUTS", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Omitted from output",
         "Notes": "0 non-null values in canonical Core; PE/PB/EV/EVEBITDA therefore impossible and not shown."},
        {"Company_ID": "(workbook)", "Ticker": "", "Field": "ROE_Calc",
         "Issue": "EQUITY_NONPOSITIVE_BLANK", "Source_Attempted": "n/a", "Retrieval_Date": TODAY,
         "Resolution": "Blank where Book_Equity missing or <= 0",
         "Notes": f"{roe_blank_eq} rows had non-positive equity and were left blank."},
    ]
    dq_new = pd.DataFrame(new_dq)
    dq_all = pd.concat([dq_old.astype(str), dq_new.astype(str)], ignore_index=True)

    readme_lines = [
        "Sector_Financials_Phase1.xlsx",
        f"Generated: {TODAY}",
        "Source: NA_Company_Financials.xlsx (canonical) - Universe (720x40), Core_Financials (720x114), Data_Quality (74)",
        "Builder: scripts/build_sector_phase1.py (pandas + openpyxl; no network access; no scraping)",
        "",
        "PURPOSE AND DISCLAIMER",
        "This is a reshape of already-collected data into one row per company for sector review.",
        "It is not a ranking, not a scorecard, and not investment advice. No buy/sell/best language appears anywhere.",
        "",
        "IDENTITY AND JOIN",
        "- Permanent key: Company_ID from Universe. Scheme: US:<TICKER>:US (500 names), CA:<TICKER>:TSX (220 names).",
        "- Join: Universe.Company_ID == Core_Financials column B (pandas name Company_ID.1), the only Core ID column using the same scheme.",
        "- Repair: CA:EFX:TSX (Enerflex Ltd.) had a blank Core column B value; its other four Core ID columns all read CA:EFX:TSX, matching Universe, so the key was backfilled from Universe and the row kept with its real financials.",
        "- Legacy Core ID columns dropped (exchange-suffix style such as US:XOM:NYSE): physical columns A, .2, .3, .4.",
        "- The seven same-ticker cross-exchange pairs remain 14 separate rows; see 04_Collisions. Ticker alone is never a key.",
        "",
        "CURRENCY WARNING",
        "Every money figure stays in its native currency (CAD for TSX names, USD for US names). No FX conversion exists in this file.",
        "Do not compare raw CAD against raw USD across borders. Check the Currency column first.",
        "",
        "WHAT COMPLETE MEANS",
        "Extraction_Status COMPLETE means the raw statement fields were pulled. It does not mean ratios exist.",
        "Blank means not available in the source data. Nothing is imputed or invented.",
        "MISSING_SOURCE rows (XOM, HONA/Honeywell Aerospace, IIP.UN) appear with blank financials on purpose.",
        "",
        "DERIVED FIELD METHODS",
        "- FCF_Reported: verbatim copy of Core.Free_Cash_Flow.",
        f"- FCF_Calc = Operating_Cash_Flow - abs(Capex). Stored Capex mixes sign conventions ({n_pos_capex} spend-positive, {n_neg_capex} outflow-negative). On the {n_div} negative-capex rows stored FCF used signed subtraction and can exceed OCF itself (ENB example: OCF 12.27B, Capex -9.165B, stored FCF 21.435B).",
        "- NetDebt_Calc = Total_Debt - Cash_ST_Investments. For Banks rows this is gross debt minus cash only, NOT a bank capital metric.",
        "- FCFMargin_Calc = FCF_Calc / Revenue. Blank for Banks.",
        "- GrossMargin_Calc = Gross_Profit / Revenue, only where GrossProfit already existed. Core has no COGS column anywhere, so no alternative gross-profit construction was used.",
        f"- {ROE_LABEL}. Blank when equity is missing or <= 0 ({roe_blank_eq} rows). Buyback-shrunken equity can produce extreme ROE (Apple ~152%); that is arithmetic, not peer-comparable quality.",
        "- Calc_Method_Notes records which of the above were applied to each row.",
        f"- Banks policy: FCF_Calc and FCFMargin_Calc are intentionally blank for the {int(banks.sum())} Banks rows.",
        "",
        "OMITTED BY DESIGN",
        "Price, Shares_Diluted and Market_Cap have zero populated values in canonical Core, so PE, PB, EV, EV/EBITDA and market cap columns are omitted entirely rather than shipped empty.",
        "",
        "SHEET MAP",
        "00_README this page; 01_All_Companies all 720 rows; 02_Coverage field fill counts; 03_Data_Quality copied canonical DQ plus new Phase 1 rows; 04_Collisions the 14 split rows; then one sheet per Custom_Industry_Sheet (30 sheets, sorted by Company_Name inside each).",
    ]
    readme = pd.DataFrame({"A": readme_lines})

    sector_names = sorted(out["Custom_Industry_Sheet"].unique())
    sheets = [("00_README", readme), ("01_All_Companies", out), ("02_Coverage", cov),
              ("03_Data_Quality", dq_all), ("04_Collisions", coll)]
    for sname in sector_names:
        sub = out[out["Custom_Industry_Sheet"].eq(sname)].sort_values(
            "Company_Name", key=lambda x: x.str.lower())
        assert sub["Company_ID"].notna().all() and sub["Company_ID"].ne("").all()
        sheets.append((sname, sub))

    if os.path.exists(OUT):
        os.remove(OUT)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        for sname, df in sheets:
            df.to_excel(xw, sheet_name=sname[:31], index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    MONEY = set(STMT_COLS) | {"FCF_Reported", "FCF_Calc", "NetDebt_Calc"}
    PCT = {"FCFMargin_Calc", "GrossMargin_Calc", "ROE_Calc"}
    wb = load_workbook(OUT)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    for sname in wb.sheetnames:
        ws = wb[sname]
        if sname == "00_README":
            ws.column_dimensions["A"].width = 150
            for row in ws.iter_rows(min_col=1, max_col=1):
                for c in row:
                    if c.value:
                        c.font = Font(bold=c.value.endswith(":") or c.value.startswith("Sector_"))
            continue
        headers = [c.value for c in ws[1]]
        for j, h in enumerate(headers, start=1):
            letter = get_column_letter(j)
            maxlen = max([len(str(h))] + [len(str(ws.cell(row=r, column=j).value)) for r in range(2, min(ws.max_row, 400) + 1) if ws.cell(row=r, column=j).value is not None], default=10)
            ws.column_dimensions[letter].width = min(46, max(11, maxlen + 2))
            hc = ws.cell(row=1, column=j)
            hc.font = hdr_font
            hc.fill = hdr_fill
            fmt = "#,##0" if h in MONEY else ("#,##0.00" if h == "Diluted_EPS" else ("0.00%" if h in PCT else None))
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=j)
                if fmt and isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(OUT)

    z = zipfile.ZipFile(OUT)
    tables = [n for n in z.namelist() if "tables/" in n.lower()]
    assert not tables, f"ListObject parts found: {tables}"
    z.close()

    wb2 = load_workbook(OUT)
    for sname in wb2.sheetnames:
        ws = wb2[sname]
        assert len(getattr(ws, "tables", {})) == 0, f"table object on {sname}"

    from python_calamine import CalamineWorkbook
    cw = CalamineWorkbook.from_path(OUT)
    rows = cw.get_sheet_by_name("01_All_Companies").to_python(skip_empty_area=False)
    hdr = rows[0]
    ix = {h: i for i, h in enumerate(hdr)}
    body = rows[1:]
    ids = [r[ix["Company_ID"]] for r in body]
    def val(cid, col):
        return next(r[ix[col]] for r in body if r[ix["Company_ID"]] == cid)
    checks = {
        "row_count_720": len(body) == 720,
        "unique_ids_720": len(set(ids)) == 720,
        "no_empty_ids": all(str(x).strip() != "" for x in ids),
        "na_name": val("CA:NA:TSX", "Company_Name") == "National Bank of Canada",
        "efx_ca_present": "CA:EFX:TSX" in ids and float(val("CA:EFX:TSX", "Revenue")) == 2571000000.0,
        "efx_us_separate": "US:EFX:US" in ids and float(val("US:EFX:US", "Revenue")) == 6074500000.0,
        "jpm_rev_numeric": isinstance(val("US:JPM:US", "Revenue"), (int, float)),
        "ry_rev_numeric": isinstance(val("CA:RY:TSX", "Revenue"), (int, float)) and float(val("CA:RY:TSX", "Revenue")) > 0,
        "hydro_one_present": "CA:H:TSX" in ids,
        "hona_present": "US:HONA:US" in ids,
        "collisions_14": sum(i in ids for i in coll_ids) == 14,
    }
    failed = [k for k, v in checks.items() if not v]
    assert not failed, f"calamine verification failed: {failed}"

    pdf = pd.read_excel(OUT, sheet_name="01_All_Companies", keep_default_na=False)
    assert len(pdf) == 720 and pdf["Company_ID"].nunique() == 720
    assert pdf["Company_ID"].isna().sum() == 0
    for sname, _ in sheets[5:]:
        sub = pd.read_excel(OUT, sheet_name=sname, keep_default_na=False)
        assert sub["Company_ID"].ne("").all(), f"empty Company_ID on {sname}"
    total_sector_rows = sum(int(out["Custom_Industry_Sheet"].eq(s).sum()) for s in sector_names)
    assert total_sector_rows == 720, f"sector partition sums to {total_sector_rows}, not 720"

    print("\n=== VERIFICATION PASSED ===")
    for k, v in checks.items():
        print(f"  {k}: {v}")

    jpm_rev = float(val("US:JPM:US", "Revenue"))
    ry_rev = float(val("CA:RY:TSX", "Revenue"))
    na_name = str(val("CA:NA:TSX", "Company_Name"))
    hydro = out[out["Company_ID"].eq("CA:H:TSX")].iloc[0]
    efx = out[out["Company_ID"].eq("CA:EFX:TSX")].iloc[0]

    rep = []
    rep.append("# Phase 1 Report")
    rep.append(f"- Date: {TODAY}")
    rep.append(f"- Output file: {OUT} ({os.path.getsize(OUT):,} bytes)")
    rep.append("- Join key used: Universe.Company_ID == Core_Financials['Company_ID.1'] (physical column B), the only Core ID column sharing the Universe scheme.")
    rep.append("- Dropped Core identity columns: Company_ID (physical col A), Company_ID.2, Company_ID.3, Company_ID.4 - all exchange-suffix style (US:XOM:NYSE) contradicting the permanent key.")
    rep.append(f"- CA:EFX:TSX repair: Core column B blank; backfilled from Universe.Company_ID (four sibling Core ID columns agreed); row retained, Revenue {float(efx['Revenue']):,.0f} CAD, Extraction_Status {efx['Extraction_Status']}.")
    rep.append(f"- All_Companies row count: 720 (unique IDs: {out['Company_ID'].nunique()}, blank IDs: {int(out['Company_ID'].eq('').sum())})")
    rep.append(f"- CA:NA:TSX = {na_name}")
    rep.append(f"- Hydro One: Company_ID {hydro['Company_ID']}, sheet {hydro['Custom_Industry_Sheet']}, Revenue {float(hydro['Revenue']):,.0f} CAD, GICS {hydro['GICS_Sector']} / {hydro['GICS_Industry']}")
    rep.append(f"- JPM Revenue: {jpm_rev:,.0f} USD | RY Revenue: {ry_rev:,.0f} CAD")
    rep.append(f"- Collisions: {len(COLLISION_TICKERS)} pairs / 14 rows retained split:")
    for t in COLLISION_TICKERS:
        us = out[out['Company_ID'].eq(f'US:{t}:US')].iloc[0]
        ca = out[out['Company_ID'].eq(f'CA:{t}:TSX')].iloc[0]
        rep.append(f"    {t}: {us['Company_Name']} (US) vs {ca['Company_Name']} (CA)")
    rep.append("")
    rep.append("Sheets:")
    rep.append("| Sheet | Rows |")
    rep.append("|---|---|")
    for sname, df in sheets:
        rep.append(f"| {sname} | {len(df)} |")
    rep.append("")
    rep.append("Coverage (of 720):")
    rep.append("| Field | n_non_null | pct | Notes |")
    rep.append("|---|---|---|---|")
    for _, r in cov.iterrows():
        rep.append(f"| {r['Field']} | {r['n_non_null']} | {r['pct']}% | {r['notes']} |")
    rep.append("")
    rep.append(f"- Remaining blank Company_ID count: {int(pdf['Company_ID'].eq('').sum())}")
    rep.append("- Confirmed: no Excel Tables/ListObjects; refresh.py not run; no scraping/network calls; no ranking or scoring; native currency preserved; Sector_Financial_Analysis.xlsx and Sector_Financials_Clean.xlsx were not used as number sources.")
    rep.append("- Second-reader verification: python-calamine reload passed all checks; openpyxl reload confirmed zero table parts.")

    with open(os.path.join(LOGD, "phase1_report.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(rep) + "\n")

    changelog = [
        "# Phase 1 Changelog",
        f"Date: {TODAY}",
        "",
        "## What changed",
        "- Created scripts/build_sector_phase1.py: Phase 1 builder joining Universe to Core_Financials on repaired Company_ID.1, computing FCF_Calc/NetDebt_Calc/FCFMargin_Calc/GrossMargin_Calc/ROE_Calc with method labels.",
        f"- Created Sector_Financials_Phase1.xlsx: {len(sheets)} sheets, 720 companies, no Excel Tables, no padding.",
        "- Created logs/phase1_report.md and this changelog.",
        "- In-file only repairs: CA:EFX:TSX join-key backfill (documented in 03_Data_Quality); nothing written back to canonical books.",
        "",
        "## Why",
        "Phase 0 proved the shipped Clean book lost Enerflex to a blank duplicate-ID join key and that derived metrics never left pilot scope. This phase rebuilds the full 720-company product from the canonical numbers with documented derived-field methods.",
        "",
        "## Files touched",
        "scripts/build_sector_phase1.py (new), Sector_Financials_Phase1.xlsx (new), logs/phase1_report.md (new), logs/phase1_changelog.md (new).",
        "",
        "## Files NOT touched",
        "NA_Company_Financials.xlsx, NA_Company_Financials_Analysis.xlsx, Sector_Financial_Analysis.xlsx, Sector_Financials_Clean.xlsx, all logs/phase0_* artifacts and audit CSV/JSON artifacts, METHODOLOGY.md, HANDOFF_STATUS.txt, all other scripts.",
        "",
        "## Known deviations from brief labels",
        "- Universe IDs are US:<TICKER>:US for all US names (brief examples showed :NASDAQ/:NYSE styles that only exist in Core/DQ legacy columns).",
        "- HONA in Universe is named Honeywell Aerospace with ID US:HONA:US; it remains the MISSING_SOURCE row distinct from Hydro One (CA:H:TSX).",
        "- NEE has blank Capex/FCF in canonical Core, so the sign convention was validated on AAPL (positive-spend) and the 212 negative-capex rows instead; documented in README and 03_Data_Quality.",
    ]
    with open(os.path.join(LOGD, "phase1_changelog.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(changelog) + "\n")

    print("\n=== TASK 3 REPORT ===")
    print("\n".join(rep))


if __name__ == "__main__":
    try:
        build()
    except SystemExit:
        raise
    except Exception:
        traceback.print_exc()
        if os.path.exists(OUT):
            try:
                os.remove(OUT)
            except OSError:
                pass
        sys.exit(1)
