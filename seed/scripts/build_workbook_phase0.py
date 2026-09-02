# -*- coding: utf-8 -*-
"""
Phase 0 — Build the empty, audit-ready workbook scaffold for the NA financials project.
Creates all always-on sheets + 11 GICS sector sheets with headers, freeze panes,
autofilters, number formats, and header comments. No financial data yet.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # na_financials_research/
OUT = os.path.join(ROOT, "NA_Company_Financials.xlsx")

# ----------------------------------------------------------------------------
# number formats
# ----------------------------------------------------------------------------
CUR   = '#,##0;(#,##0)'          # currency, thousands, negatives in parens
CUR2  = '#,##0.00;(#,##0.00)'    # currency, 2dp (price, EPS, per-share)
PCT   = '0.0%'
MULT  = '0.0"x"'                 # valuation multiples
RATIO = '0.00'
RATIO3= '0.000'
INT   = '0'
NUM   = '0.0'

# ----------------------------------------------------------------------------
# Core_Financials column definitions
# (name, fmt, direction, unit, definition, formula)
# ----------------------------------------------------------------------------
CORE = [
    # identity / market
    ("Ticker", None, "neutral", "text", "Primary ticker (US: AAPL; CA: RY.TO)", ""),
    ("Company_Name", None, "neutral", "text", "Legal / common name", ""),
    ("Custom_Industry_Sheet", None, "neutral", "text", "Mapping to industry detail sheet", ""),
    ("GICS_Sector", None, "neutral", "text", "GICS sector", ""),
    ("Price", CUR2, "neutral", "native currency", "Latest price in native reporting currency", ""),
    ("Market_Cap", CUR, "neutral", "native currency", "Price x shares diluted", "Price*Shares_Diluted"),
    ("Enterprise_Value", CUR, "neutral", "native currency", "MktCap + debt + preferred + minority - cash & ST inv", "MktCap+Debt+Pref+Minority-Cash"),
    ("Shares_Diluted", CUR, "neutral", "shares", "Diluted shares outstanding", ""),
    ("Dividend_Yield", PCT, "context", "percent", "DPS / Price", "DPS/Price"),
    ("Payout_Ratio", PCT, "context", "percent", "Dividends / Net income", "Div/NI"),
    ("Buyback_Yield", PCT, "higher_better", "percent", "-change in diluted shares (net issuance = negative)", "-dShares/Shares_prior"),
    # income statement
    ("Revenue", CUR, "neutral", "native currency", "Total revenue (TTM or latest annual)", ""),
    ("Gross_Profit", CUR, "neutral", "native currency", "Gross profit", ""),
    ("Gross_Margin", PCT, "higher_better", "percent", "Gross profit / revenue (pricing power)", "GP/Rev"),
    ("Operating_Income_EBIT", CUR, "neutral", "native currency", "Operating income (EBIT)", ""),
    ("Operating_Margin", PCT, "higher_better", "percent", "EBIT / revenue", "EBIT/Rev"),
    ("EBITDA", CUR, "neutral", "native currency", "EBIT + D&A", "EBIT+D&A"),
    ("EBITDA_Adj", CUR, "neutral", "native currency", "Company-adjusted EBITDA (if differs from reported)", ""),
    ("Net_Income", CUR, "neutral", "native currency", "Net income to common", ""),
    ("Diluted_EPS", CUR2, "higher_better", "native currency", "Diluted earnings per share", ""),
    ("Interest_Expense", CUR, "neutral", "native currency", "Interest expense", ""),
    ("Tax_Rate_Effective", PCT, "context", "percent", "Tax / pretax income (cap 35% in NOPAT unless statutory)", "Tax/Pretax"),
    ("R_D", CUR, "neutral", "native currency", "Research & development spend (if disclosed)", ""),
    ("SG_A", CUR, "neutral", "native currency", "Selling, general & admin (if disclosed)", ""),
    ("Revenue_CAGR_3y", PCT, "higher_better", "percent", "3-year revenue CAGR", "(Rev/Rev-3)^(1/3)-1"),
    ("Revenue_CAGR_5y", PCT, "higher_better", "percent", "5-year revenue CAGR", "(Rev/Rev-5)^(1/5)-1"),
    ("EPS_CAGR_3y", PCT, "higher_better", "percent", "3-year diluted EPS CAGR", "(EPS/EPS-3)^(1/3)-1"),
    ("EPS_CAGR_5y", PCT, "higher_better", "percent", "5-year diluted EPS CAGR", "(EPS/EPS-5)^(1/5)-1"),
    ("Share_Count_CAGR_5y", PCT, "lower_better", "percent", "5-year share-count CAGR (flags buyback-driven EPS growth)", "(Sh/Sh-5)^(1/5)-1"),
    # balance sheet / liquidity / leverage
    ("Cash_ST_Investments", CUR, "neutral", "native currency", "Cash + short-term investments", ""),
    ("Receivables", CUR, "neutral", "native currency", "Accounts receivable", ""),
    ("Inventory", CUR, "neutral", "native currency", "Inventory", ""),
    ("Payables", CUR, "neutral", "native currency", "Accounts payable", ""),
    ("Current_Assets", CUR, "neutral", "native currency", "Total current assets", ""),
    ("Current_Liabilities", CUR, "neutral", "native currency", "Total current liabilities", ""),
    ("Total_Assets", CUR, "neutral", "native currency", "Total assets", ""),
    ("Total_Debt", CUR, "neutral", "native currency", "Short + long interest-bearing debt", ""),
    ("Lease_Liabilities", CUR, "neutral", "native currency", "IFRS 16 / ASC 842 lease liabilities (separate)", ""),
    ("Net_Debt", CUR, "lower_better", "native currency", "Total debt - cash & ST inv", "Debt-Cash"),
    ("Book_Equity", CUR, "neutral", "native currency", "Total book equity", ""),
    ("Tangible_Book", CUR, "neutral", "native currency", "Equity - goodwill - other intangibles", "Eq-GW-Intang"),
    ("Current_Ratio", RATIO, "context", "ratio", "Current assets / current liabilities (negative WC ok for Amazon-style)", "CA/CL"),
    ("Quick_Ratio", RATIO, "context", "ratio", "(Cash + receivables) / current liabilities", "(Cash+AR)/CL"),
    ("Net_Working_Capital", CUR, "context", "native currency", "Current assets - current liabilities", "CA-CL"),
    ("Debt_Equity", RATIO, "lower_better", "ratio", "Total debt / book equity (NOT for banks)", "Debt/Eq"),
    ("Net_Debt_EBITDA", RATIO, "lower_better", "ratio", "Net debt / EBITDA (NOT for banks/insurers)", "NetDebt/EBITDA"),
    ("Interest_Coverage", RATIO, "higher_better", "ratio", "EBIT / interest (flag net vs capitalized)", "EBIT/Int"),
    ("Cash_Conversion_Cycle", NUM, "lower_better", "days", "DSO + DIO - DPO (not for banks)", "DSO+DIO-DPO"),
    # cash flow / earnings quality
    ("Operating_Cash_Flow", CUR, "neutral", "native currency", "Operating cash flow", ""),
    ("Capex", CUR, "neutral", "native currency", "Capital expenditures (reported)", ""),
    ("Free_Cash_Flow", CUR, "higher_better", "native currency", "OCF - capex (default reported capex)", "OCF-Capex"),
    ("FCF_Margin", PCT, "higher_better", "percent", "FCF / revenue", "FCF/Rev"),
    ("FCF_NI", RATIO, "higher_better", "ratio", "FCF / net income (cash conversion)", "FCF/NI"),
    ("Owner_Earnings", CUR, "higher_better", "native currency", "NI + D&A - maintenance capex (use capex + flag if maintenance unknown)", "NI+D&A-mCapex"),
    ("Accruals", RATIO3, "lower_better", "ratio", "(NI - OCF) / total assets (Sloan-style)", "(NI-OCF)/TA"),
    ("Capex_D_A", RATIO, "context", "ratio", "Capex / D&A (reinvestment intensity)", "Capex/D&A"),
    # returns (quality)
    ("ROE", PCT, "higher_better", "percent", "Net income / average equity (DuPont)", "NI/avgEq"),
    ("ROA", PCT, "higher_better", "percent", "Net income / average assets", "NI/avgTA"),
    ("ROIC", PCT, "higher_better", "percent", "NOPAT / invested capital (best industrial quality metric)", "NOPAT/IC"),
    ("NOPAT", CUR, "neutral", "native currency", "EBIT x (1 - tax rate), effective tax floor 0 cap 35%", "EBIT*(1-t)"),
    ("Invested_Capital", CUR, "neutral", "native currency", "Equity + debt + leases - cash (avg begin/end)", "Eq+Debt+Lease-Cash"),
    ("ROIC_WACC_Spread", PCT, "higher_better", "percent", "ROIC - Damodaran industry WACC (estimate)", "ROIC-WACC"),
    ("Gross_Profitability", PCT, "higher_better", "percent", "Gross profit / total assets (Novy-Marx)", "GP/TA"),
    ("DuPont_Net_Margin", PCT, "higher_better", "percent", "Net income / revenue", "NI/Rev"),
    ("DuPont_Asset_Turnover", RATIO, "higher_better", "ratio", "Revenue / average assets", "Rev/avgTA"),
    ("DuPont_Equity_Multiplier", RATIO, "lower_better", "ratio", "Average assets / average equity (leverage)", "avgTA/avgEq"),
    # valuation
    ("PE_Trailing", MULT, "lower_better", "multiple", "Price / trailing EPS", "P/EPS"),
    ("PE_Forward", MULT, "lower_better", "multiple", "Price / forward consensus EPS (if available)", "P/EPS_fwd"),
    ("PEG", RATIO, "lower_better", "ratio", "P/E / 5y EPS or revenue CAGR (NA if growth <= 0)", "PE/CAGR"),
    ("PB", MULT, "lower_better", "multiple", "Price / book value per share", "P/BVPS"),
    ("P_Tangible_Book", MULT, "lower_better", "multiple", "Price / tangible book per share", "P/TBVPS"),
    ("EV_EBIT", MULT, "lower_better", "multiple", "Enterprise value / EBIT", "EV/EBIT"),
    ("EV_EBITDA", MULT, "lower_better", "multiple", "Enterprise value / EBITDA", "EV/EBITDA"),
    ("EV_Sales", MULT, "lower_better", "multiple", "Enterprise value / revenue", "EV/Rev"),
    ("Earnings_Yield", PCT, "higher_better", "percent", "EBIT / EV (Greenblatt definition)", "EBIT/EV"),
    ("FCF_Yield_MktCap", PCT, "higher_better", "percent", "FCF / market cap", "FCF/MktCap"),
    ("FCF_Yield_EV", PCT, "higher_better", "percent", "FCF / EV", "FCF/EV"),
    ("Graham_Number", CUR2, "context", "native currency", "sqrt(22.5 x EPS x BVPS) - conservative reference only", "sqrt(22.5*EPS*BVPS)"),
    ("Div_Buyback_Yield", PCT, "higher_better", "percent", "Dividend yield + buyback yield", "DivYield+BuybackYield"),
    # composite academic scores
    ("Piotroski_F_Score", INT, "higher_better", "0-9", "Piotroski F-Score (7-9 strong, 4-6 mixed, 0-3 weak). Not for banks.", ""),
    ("Altman_Z_Score", RATIO, "higher_better", "ratio", "Altman Z (>2.99 safe, 1.81-2.99 grey, <1.81 distress). Not for banks.", ""),
    ("Altman_Z_Variant", None, "neutral", "text", "Z (manufacturers) / Z' / Z'' variant used", ""),
    ("Beneish_M_Score", RATIO, "lower_better", "ratio", "Beneish M-Score; flag if > -2.22", ""),
    ("Greenblatt_ROC", PCT, "higher_better", "percent", "EBIT / (NWC + net fixed assets); NWC floored at 0", "EBIT/(NWC+NFA)"),
    ("Greenblatt_Earnings_Yield", PCT, "higher_better", "percent", "EBIT / EV", "EBIT/EV"),
    ("Greenblatt_Combined_Rank", INT, "lower_better", "rank", "rank(ROC) + rank(EY); lower better", ""),
    ("Buffett_Checklist_Count", INT, "higher_better", "0-6", "Buffett-style quality checklist count (0-6)", ""),
    # source tags
    ("Source_Primary", None, "neutral", "text", "Filing type + period (e.g. '10-K FY2024')", ""),
    ("Source_Aggregator", None, "neutral", "text", "Aggregator name (e.g. 'Yahoo Finance')", ""),
    ("Retrieval_Date", None, "neutral", "date", "Date the number was retrieved", ""),
    ("Notes", None, "neutral", "text", "Free notes / flags", ""),
]

# ----------------------------------------------------------------------------
# Sector-specific metric definitions  (name, sheet, direction, unit, definition)
# ----------------------------------------------------------------------------
SECTOR_METRICS = [
    # Banks
    ("Net_Interest_Margin", "Banks", "higher_better", "percent", "Net interest income / average earning assets"),
    ("NII_Total_Revenue", "Banks", "context", "percent", "Net interest income / total revenue"),
    ("Efficiency_Ratio", "Banks", "lower_better", "percent", "Non-interest expense / revenue (lower better)"),
    ("CET1_Ratio", "Banks", "higher_better", "percent", "Common Equity Tier 1 capital ratio"),
    ("Total_Capital_Ratio", "Banks", "higher_better", "percent", "Total regulatory capital ratio"),
    ("Liquidity_Coverage_Ratio", "Banks", "higher_better", "percent", "LCR if disclosed"),
    ("Loan_Deposit_Ratio", "Banks", "context", "percent", "Loans / deposits"),
    ("NPL_Gross_Loans", "Banks", "lower_better", "percent", "Non-performing loans / gross loans (impaired loans ratio)"),
    ("Net_Charge_Off_Ratio", "Banks", "lower_better", "percent", "Net charge-offs / average loans"),
    ("Allowance_NPL_Coverage", "Banks", "higher_better", "percent", "Allowance / NPLs (coverage)"),
    ("Avg_Loans", "Banks", "neutral", "native currency", "Average loans"),
    ("Avg_Deposits", "Banks", "neutral", "native currency", "Average deposits"),
    ("Book_Value_Share", "Banks", "higher_better", "native currency", "Book value per share"),
    ("Tangible_Book_Share", "Banks", "higher_better", "native currency", "Tangible book value per share"),
    ("Provision_Credit_Losses", "Banks", "neutral", "native currency", "Provision for credit losses"),
    # Insurance
    ("Gross_Premiums_Written", "Insurance", "higher_better", "native currency", "Gross premiums written"),
    ("Net_Premiums_Earned", "Insurance", "higher_better", "native currency", "Net premiums earned"),
    ("Combined_Ratio", "Insurance", "lower_better", "percent", "P&C combined ratio (lower better)"),
    ("Loss_Ratio", "Insurance", "lower_better", "percent", "Loss ratio (P&C)"),
    ("Expense_Ratio", "Insurance", "lower_better", "percent", "Expense ratio (P&C)"),
    ("Investment_Yield", "Insurance", "higher_better", "percent", "Investment yield"),
    ("Float_Proxy", "Insurance", "context", "native currency", "Float proxy if computable"),
    ("RBC_or_MCT", "Insurance", "higher_better", "percent", "RBC (US) / MCT (Canada) solvency"),
    ("Reserve_Development", "Insurance", "context", "native currency", "Reserve development if disclosed"),
    ("Benefit_Ratio", "Insurance", "lower_better", "percent", "Life: benefit ratio"),
    ("New_Business_CSM", "Insurance", "higher_better", "native currency", "Life: new business CSM (IFRS 17)"),
    # Credit services
    ("Net_Charge_Offs", "Credit_Services", "lower_better", "percent", "Net charge-offs"),
    ("Delinquency_30_90", "Credit_Services", "lower_better", "percent", "30/90 day delinquency"),
    ("Credit_Loss_Reserve_Loans", "Credit_Services", "higher_better", "percent", "Credit loss reserve / loans"),
    # Software / SaaS
    ("Rule_of_40", "Software", "higher_better", "percent", "Revenue growth % + FCF margin %"),
    ("RPO_Deferred_Revenue", "Software", "higher_better", "native currency", "Remaining performance obligations / deferred revenue"),
    ("R_D_Pct_Sales", "Software", "context", "percent", "R&D / sales"),
    ("S_M_Pct_Sales", "Software", "lower_better", "percent", "S&M / sales"),
    ("Net_Revenue_Retention", "Software", "higher_better", "percent", "NRR if disclosed; else NA"),
    ("cRPO", "Software", "higher_better", "native currency", "Current remaining performance obligations"),
    # Cloud
    ("Cloud_Revenue", "Cloud", "higher_better", "native currency", "Cloud segment revenue (from segment note)"),
    ("Cloud_Growth", "Cloud", "higher_better", "percent", "Cloud segment revenue growth"),
    ("Cloud_Operating_Margin", "Cloud", "higher_better", "percent", "Cloud operating margin if disclosed"),
    ("Capex_Sales", "Cloud", "context", "percent", "Capex / sales"),
    # Semiconductors / components
    ("Inventory_Days", "Semiconductors_Components", "lower_better", "days", "Inventory days"),
    ("Book_To_Bill", "Semiconductors_Components", "higher_better", "ratio", "Book-to-bill if disclosed"),
    # Networking
    ("Backlog", "Networking", "higher_better", "native currency", "Backlog if disclosed"),
    # Internet platforms
    ("MAU_DAU", "Internet_Platforms", "higher_better", "count", "MAU/DAU or subscribers if reported"),
    ("Take_Rate_ARPU", "Internet_Platforms", "higher_better", "native currency", "Take rate / ARPU if reported"),
    ("SBC_Revenue", "Internet_Platforms", "lower_better", "percent", "Stock-based compensation / revenue (quality drag)"),
    # Telecom
    ("Subscribers", "Telecom", "higher_better", "count", "Subscribers by segment if reported"),
    ("ARPU", "Telecom", "higher_better", "native currency", "Average revenue per user"),
    ("Churn", "Telecom", "lower_better", "percent", "Churn if reported"),
    ("Dividend_Coverage_FCF", "Telecom", "higher_better", "ratio", "Dividend coverage by FCF"),
    # Streaming / entertainment
    ("Paid_Subscribers", "Streaming_Entertainment", "higher_better", "count", "Paid subscribers"),
    ("Content_Spend", "Streaming_Entertainment", "context", "native currency", "Content spend (capitalization differences footnoted)"),
    # Retail / discount / consumer goods
    ("Same_Store_Sales", "Retail", "higher_better", "percent", "Same-store / comparable sales if disclosed"),
    ("SG_A_Pct", "Retail", "lower_better", "percent", "SG&A / sales"),
    ("Inventory_Turnover_DIO", "Retail", "lower_better", "days", "Inventory turnover / DIO (lower DIO better)"),
    ("Membership_Income", "Retail", "higher_better", "native currency", "Membership income if any (Costco/Walmart/Dollarama)"),
    ("Sales_Per_SqFt", "Retail", "higher_better", "native currency", "Sales per square foot if disclosed"),
    # Fast food / restaurants
    ("Systemwide_Sales", "Fast_Food_Restaurants", "higher_better", "native currency", "Systemwide sales"),
    ("Franchise_Company_Mix", "Fast_Food_Restaurants", "context", "percent", "Franchise vs company mix"),
    ("Restaurant_Level_Margin", "Fast_Food_Restaurants", "higher_better", "percent", "Restaurant-level margin if disclosed"),
    ("Unit_Growth", "Fast_Food_Restaurants", "higher_better", "percent", "Unit growth"),
    ("Net_New_Units", "Fast_Food_Restaurants", "higher_better", "count", "Net new units"),
    # Pharma
    ("Pipeline_Comments", "Pharma", "context", "text", "Cash/pipeline comments from filing only"),
    ("LOE_Flags", "Pharma", "context", "text", "Loss-of-exclusivity flags from 10-K risk factors"),
    # Biotech
    ("Cash_Runway", "Biotech", "higher_better", "months", "Cash / TTM operating cash burn (if unprofitable)"),
    ("R_D_Spend", "Biotech", "context", "native currency", "R&D spend"),
    ("Dilution_History", "Biotech", "context", "text", "Dilution history"),
    # Medical devices / services / managed care
    ("Medical_Loss_Ratio", "Medical_Devices_Services", "lower_better", "percent", "Managed care: medical loss ratio"),
    ("Occupancy", "Medical_Devices_Services", "higher_better", "percent", "Facilities: occupancy if disclosed"),
    # Railroads
    ("Operating_Ratio", "Railroads", "lower_better", "percent", "Operating expense / revenue (lower better)"),
    ("Volume_RTM", "Railroads", "higher_better", "count", "Volume / revenue ton-miles if disclosed"),
    # Airlines
    ("Load_Factor", "Airlines", "higher_better", "percent", "Load factor"),
    ("CASM", "Airlines", "lower_better", "native currency", "Cost per available seat mile"),
    ("RASM", "Airlines", "higher_better", "native currency", "Revenue per available seat mile"),
    ("Available_Seat_Miles", "Airlines", "higher_better", "count", "Available seat miles"),
    ("Liquidity", "Airlines", "higher_better", "native currency", "Cash + revolver liquidity"),
    # Autos
    ("Auto_Operating_Margin", "Autos", "higher_better", "percent", "Automotive operating margin (store adjusted separately)"),
    ("Net_Cash_Debt", "Autos", "higher_better", "native currency", "Net cash (positive) / net debt (negative)"),
    ("Warranty_Recall_Notes", "Autos", "context", "text", "Warranty / recall notes only if quantified"),
    # Energy producers
    ("Production_Volume", "Oil_Gas_Producers", "higher_better", "count", "Production volumes if disclosed"),
    ("Realized_Price", "Oil_Gas_Producers", "context", "native currency", "Realized price if disclosed"),
    ("Capex_vs_D_A", "Oil_Gas_Producers", "context", "ratio", "Capex vs D&A"),
    ("Reserve_Data", "Oil_Gas_Producers", "context", "text", "Reserve data only if in filing; else NA"),
    # Midstream / pipelines
    ("Distributable_Cash_Flow", "Pipelines_Midstream", "higher_better", "native currency", "Distributable cash flow or FCF"),
    ("Distribution_Coverage", "Pipelines_Midstream", "higher_better", "ratio", "Distribution / dividend coverage"),
    ("Debt_EBITDA", "Pipelines_Midstream", "lower_better", "ratio", "Debt / EBITDA"),
    ("Contracted_Volume_Pct", "Pipelines_Midstream", "higher_better", "percent", "Contracted volume % if disclosed"),
    # Utilities (regulated)
    ("Allowed_Or_Earned_ROE", "Utilities_Regulated", "higher_better", "percent", "Allowed/earned ROE if disclosed"),
    ("Rate_Base", "Utilities_Regulated", "higher_better", "native currency", "Rate base"),
    ("Rate_Base_Growth", "Utilities_Regulated", "higher_better", "percent", "Rate-base growth if disclosed"),
    ("FFO_Debt", "Utilities_Regulated", "higher_better", "ratio", "FFO / debt (or OCF / debt)"),
    ("Dividend_Growth", "Utilities_Regulated", "higher_better", "percent", "Dividend growth"),
    ("Regulatory_Jurisdiction", "Utilities_Regulated", "context", "text", "Regulatory jurisdiction (e.g. Ontario for Hydro One)"),
    # Real estate / REITs
    ("FFO", "Real_Estate", "higher_better", "native currency", "Funds from operations"),
    ("AFFO", "Real_Estate", "higher_better", "native currency", "Adjusted funds from operations if reported"),
    ("NAV", "Real_Estate", "context", "native currency", "NAV or stated book"),
    ("AFFO_Payout", "Real_Estate", "lower_better", "percent", "AFFO payout ratio"),
    # Materials / mining
    ("AISC", "Materials", "lower_better", "native currency", "All-in sustaining costs if disclosed"),
    ("Cash_Cost", "Materials", "lower_better", "native currency", "Cash cost if disclosed"),
    ("Reserve_Life", "Materials", "higher_better", "years", "Reserve life if disclosed"),
]

# ----------------------------------------------------------------------------
# Sector sheets (GICS) — Core columns + score/rank columns
# ----------------------------------------------------------------------------
GICS_SECTORS = [
    "Tech", "Comm_Services", "Consumer_Cyclical", "Consumer_Defensive",
    "Healthcare", "Financials", "Energy", "Industrials", "Materials",
    "Utilities", "Real_Estate",
]

SCORE_COLS = [
    ("Sector_Percentile", PCT, "higher_better", "percent", "Composite percentile within sector/industry"),
    ("Quality_Score", NUM, "higher_better", "0-100", "Quality sub-score"),
    ("Health_Score", NUM, "higher_better", "0-100", "Health sub-score"),
    ("Growth_Score", NUM, "higher_better", "0-100", "Growth sub-score"),
    ("Value_Score", NUM, "higher_better", "0-100", "Value sub-score"),
    ("Composite_Score", NUM, "higher_better", "0-100", "Composite score (weighted)"),
    ("Sector_Rank", INT, "lower_better", "rank", "Rank within sheet (1 = best)"),
    ("Red_Flags", None, "neutral", "text", "Short red-flag text"),
    ("Why_It_Ranks", None, "neutral", "text", "1-2 sentences, numbers only, <=40 words"),
]

# header comments for non-obvious core columns
COMMENTS = {
    "Accruals": "Sloan (1996): earnings that do not convert to cash. Negative (NI < OCF) is better.",
    "Owner_Earnings": "Buffett owner earnings. Uses reported capex if maintenance capex unknown — flagged.",
    "EBITDA_Adj": "Only populated when company-adjusted EBITDA differs from reported; leave blank otherwise.",
    "Buyback_Yield": "Net issuance = negative. Uses diluted share count change.",
    "Tax_Rate_Effective": "NOPAT uses effective tax, floor 0, cap 35% unless statutory disclosed.",
    "Altman_Z_Score": "Not valid for banks/insurers. Variant recorded in Altman_Z_Variant.",
    "Net_Debt_EBITDA": "Not applicable to banks/insurers — leave Data_Quality.",
    "Debt_Equity": "Not applicable to banks/insurers — leave Data_Quality.",
    "Greenblatt_ROC": "NWC floored at 0 when negative.",
    "PEG": "NA if growth <= 0.",
    "Cash_Conversion_Cycle": "Not applicable to banks.",
    "Interest_Coverage": "Flag whether interest is net or capitalized.",
}

# ----------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGN = Alignment(vertical="top")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, headers, width_map=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
        if h in COMMENTS:
            cell.comment = Comment(COMMENTS[h], "OpenClaw")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    if width_map:
        for c, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(c)].width = width_map.get(h, 14)


def apply_num_format(ws, headers, fmt_map, n_rows=1500):
    """Apply number formats to a generous data range so values show correctly when filled."""
    for c, h in enumerate(headers, start=1):
        fmt = fmt_map.get(h)
        if not fmt:
            continue
        for r in range(2, n_rows + 1):
            ws.cell(row=r, column=c).number_format = fmt


def color_scale(ws, col_letter, n_rows=1500):
    rng = f"{col_letter}2:{col_letter}{n_rows}"
    ws.conditional_formatting.add(
        rng, ColorScaleRule(start_type="min", start_color="F8696B",
                            mid_type="percentile", mid_value=50, mid_color="FFEB84",
                            end_type="max", end_color="63BE7B")
    )


def text_sheet(ws, title, rows):
    """rows = list of (value, bold) or list of strings."""
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    r = 3
    for row in rows:
        if isinstance(row, tuple):
            val, bold = row
        else:
            val, bold = row, False
        cell = ws.cell(row=r, column=1, value=val)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True)
        r += 1
    ws.column_dimensions["A"].width = 130
    ws.freeze_panes = "A2"


# ----------------------------------------------------------------------------
# build workbook
# ----------------------------------------------------------------------------
wb = openpyxl.Workbook()
wb.remove(wb.active)

core_names = [c[0] for c in CORE]
core_fmt = {c[0]: c[1] for c in CORE}
core_dir = {c[0]: c[2] for c in CORE}
core_unit = {c[0]: c[3] for c in CORE}
core_def  = {c[0]: c[4] for c in CORE}
core_fmla = {c[0]: c[5] for c in CORE}

# 1. README
ws = wb.create_sheet("README")
text_sheet(ws, "NA Company Financials Workbook — README", [
    ("How to use", True),
    "This workbook compares US and Canada public companies by sector using financial-statement data (10-K/10-Q/20-F/AIF/MD&A), not marketing copy. All numeric cells are source-tagged.",
    "Start at Universe (who is in scope) -> Core_Financials (shared metrics) -> Scores (rankings) -> GICS sector sheets -> industry detail sheets (Phase 4).",
    "",
    ("Last refresh", True),
    "Phase 0 — scaffold only. No financial data yet. Last refresh: (pending Phase 3)",
    "",
    ("FX rate", True),
    "USD/CAD rate + date: (pending Phase 3). Native currency kept; USD columns use documented rate.",
    "",
    ("Coverage counts", True),
    "Universe: 0 tickers | Core_Financials: 0 rows | Scores: 0 rows (Phase 0).",
    "",
    ("Rules (see Methodology sheet + METHODOLOGY.md)", True),
    "Never invent numbers; Data_Quality tag + log when unsourceable. Source hierarchy: EDGAR XBRL > SEDAR+ > IR > aggregators. Ratios are Excel formulas; values cached. No buy/sell advice — within-sector rank only.",
])

# 2. Methodology
ws = wb.create_sheet("Methodology")
meth_lines = [
    ("Scoring model (within sector/industry; percentiles 0-100, higher = better)", True),
    "Non-financial: Quality 40% (ROIC, ROIC 5y, gross-margin stability, FCF/NI, Buffett checklist, gross profitability) | Health 25% (interest coverage, NetDebt/EBITDA, current ratio, Altman Z if valid, F-Score) | Growth 15% (Rev CAGR 5y, FCF CAGR, share-count penalty) | Value 20% (earnings yield, FCF yield, EV/EBITDA vs peer, PEG).",
    "Banks: Quality 35% (ROE, ROA, efficiency ratio inverted, NIM vs peer) | Health 35% (CET1, NPL, NCOs, coverage) | Growth 10% (BVPS CAGR, loan growth) | Value 20% (P/TBV, P/E, dividend+earnings yield).",
    "Utilities: Quality 30% (ROE vs allowed, FFO/Debt, earned returns) | Health 30% (FFO/Debt, interest coverage, payout sustainability) | Growth 20% (rate-base or EPS/DPS CAGR) | Value 20% (P/E, dividend yield vs growth, P/B).",
    "",
    ("Metric definitions", True),
    "See METHODOLOGY.md and the Data_Dictionary sheet for formulas, units, good/bad direction, and sector applicability.",
    "",
    ("Academic sources", True),
    "Greenblatt ROC+EY; Novy-Marx gross profitability; AQR Quality-Minus-Junk; Sloan accruals; Piotroski F-Score; Altman Z/Z'/Z''; Beneish M-Score; Damodaran industry WACC/sector medians.",
    "",
    ("Piotroski F-Score (0-9)", True),
    "ROA>0; OCF>0; ROA up YoY; OCF>NI; LT-debt ratio down; current ratio up; no share issuance; gross margin up; asset turnover up. 7-9 strong / 4-6 mixed / 0-3 weak. Not for banks.",
    "",
    ("Buffett checklist (0-6)", True),
    "ROE>=15% in >=4/5y; ROIC>=15% latest+5y avg; positive FCF in >=4/5y; NetDebt/EBITDA<=2.5 (non-fin); operating margin not down >250bps vs 5y ago; no serial dilution.",
]
text_sheet(ws, "Methodology", meth_lines)

# 3. Universe
universe_cols = [
    "Universe_Effective_Date", "Ticker", "Primary_Ticker", "Yahoo_Ticker",
    "Secondary_Ticker", "Share_Class", "Company_Name", "Country_of_Listing",
    "Exchange", "Primary_Currency", "Financials_Native_Currency", "Trading_Currency",
    "Reporting_Standard", "GICS_Sector", "GICS_Industry_Group", "GICS_Industry",
    "GICS_Sub_Industry", "Custom_Industry_Sheet", "Market_Cap_Native",
    "Market_Cap_USD", "Shares_Out", "Price", "Price_Date", "Market_Data_As_Of",
    "Financials_As_Of", "Fiscal_Year_End", "CIK_SEDAR", "IR_URL",
    "Latest_Annual_Filing", "Latest_Interim_Filing", "Dual_Listed",
    "Include_In_Ranking", "In_SP500", "In_TSX_Composite", "In_TSX60", "In_SPUS",
    "Data_Source", "Retrieval_Date", "Notes",
]
uni_fmt = {
    "Market_Cap_Native": CUR, "Market_Cap_USD": CUR, "Shares_Out": CUR, "Price": CUR2,
}
uni_width = {c: 16 for c in universe_cols}
for c in ["Company_Name", "IR_URL", "Data_Source", "Notes", "Secondary_Ticker", "Universe_Effective_Date"]:
    uni_width[c] = 26
ws = wb.create_sheet("Universe")
style_header(ws, universe_cols, uni_width)
apply_num_format(ws, universe_cols, uni_fmt)
ws.cell(row=1, column=36).comment = Comment("In_SPUS = present in the attached SPUS (sharia-screened S&P 500) holdings list. Optional extension flag.", "OpenClaw")

# Source_Audit (long-format audit trail) — created in Phase 0, populated from Phase 2
sa_cols = [
    "Ticker", "Fiscal_Period", "Metric", "Value", "Unit", "Source_URL",
    "Filing_Type", "Filing_Date", "Page_or_XBRL_Tag", "Retrieval_Date",
    "Extraction_Method", "Confidence", "Notes",
]
ws = wb.create_sheet("Source_Audit")
style_header(ws, sa_cols, {"Metric": 24, "Source_URL": 40, "Page_or_XBRL_Tag": 24, "Notes": 30, "Extraction_Method": 18})

# 4. Data_Dictionary
dd_cols = ["Column_Name", "Sheet", "Definition", "Formula", "Unit", "Direction", "Sector_Applicability", "Notes"]
dd_width = {"Column_Name": 26, "Sheet": 20, "Definition": 60, "Formula": 34, "Unit": 14, "Direction": 14, "Sector_Applicability": 26, "Notes": 30}
ws = wb.create_sheet("Data_Dictionary")
style_header(ws, dd_cols, dd_width)
r = 2
for c in CORE:
    name, fmt, direction, unit, definition, formula = c
    applicability = "All"
    if name in ("Debt_Equity", "Net_Debt_EBITDA", "Altman_Z_Score", "Altman_Z_Variant", "Cash_Conversion_Cycle"):
        applicability = "Non-financial (not banks/insurers)"
    if name in ("Piotroski_F_Score",):
        applicability = "All except banks (do not rank banks on F-score)"
    row = [name, "Core_Financials", definition, formula, unit, direction, applicability, ""]
    for cidx, v in enumerate(row, start=1):
        ws.cell(row=r, column=cidx, value=v)
    r += 1
for m in SECTOR_METRICS:
    name, sheet, direction, unit, definition = m
    row = [name, sheet, definition, "", unit, direction, sheet, ""]
    for cidx, v in enumerate(row, start=1):
        ws.cell(row=r, column=cidx, value=v)
    r += 1
ws.auto_filter.ref = f"A1:H{r-1}"

# 5. Core_Financials
ws = wb.create_sheet("Core_Financials")
style_header(ws, core_names)
apply_num_format(ws, core_names, core_fmt)

# 6. Time_Series (long format)
ts_cols = ["Ticker", "Fiscal_Year", "Field", "Value", "Unit", "Currency", "Source", "Retrieval_Date", "Notes"]
ws = wb.create_sheet("Time_Series")
style_header(ws, ts_cols, {"Field": 24, "Source": 24, "Notes": 24, "Value": 16})

# 7. Sector_Benchmarks
sb_cols = ["GICS_Sector", "Industry", "Metric", "Median", "Mean", "n", "P25", "P75", "Source", "Retrieval_Date", "Notes"]
sb_fmt = {"Median": CUR, "Mean": CUR, "P25": CUR, "P75": CUR}
ws = wb.create_sheet("Sector_Benchmarks")
style_header(ws, sb_cols, {"Metric": 24, "Source": 24, "Notes": 24})
apply_num_format(ws, sb_cols, sb_fmt)

# 8. Scores
scores_cols = [
    "Ticker", "Company_Name", "GICS_Sector", "Custom_Industry_Sheet",
    "Quality_Score", "Health_Score", "Growth_Score", "Value_Score", "Composite_Score",
    "Sector_Rank", "Sector_Percentile", "Red_Flags", "Why_It_Ranks", "Low_Sample_Flag", "Notes",
]
scores_fmt = {"Quality_Score": NUM, "Health_Score": NUM, "Growth_Score": NUM,
              "Value_Score": NUM, "Composite_Score": NUM, "Sector_Rank": INT, "Sector_Percentile": PCT}
ws = wb.create_sheet("Scores")
style_header(ws, scores_cols, {"Company_Name": 26, "Red_Flags": 40, "Why_It_Ranks": 40, "Notes": 26})
apply_num_format(ws, scores_cols, scores_fmt)
color_scale(ws, "I")   # Composite_Score
color_scale(ws, "K")   # Sector_Percentile

# 9. Data_Quality
dq_cols = ["Ticker", "Field", "Issue", "Source_Attempted", "Retrieval_Date", "Resolution", "Notes"]
ws = wb.create_sheet("Data_Quality")
style_header(ws, dq_cols, {"Issue": 40, "Source_Attempted": 30, "Resolution": 30, "Notes": 30})

# 10. Watchlist
wl_cols = ["Ticker", "Company_Name", "Added_Date", "Rationale", "Notes"]
ws = wb.create_sheet("Watchlist")
style_header(ws, wl_cols, {"Rationale": 40, "Notes": 30, "Company_Name": 26})

# 11. Sector_Leaders
sl_cols = [
    "GICS_Sector", "Industry", "Sector_Rank", "Ticker", "Company_Name",
    "Composite_Score", "Quality_Score", "Health_Score", "Growth_Score", "Value_Score",
    "ROIC", "Gross_Margin", "FCF_Margin", "Revenue_CAGR_5y", "Earnings_Yield",
    "Net_Debt_EBITDA", "Why_It_Ranks",
]
sl_fmt = {"Composite_Score": NUM, "Quality_Score": NUM, "Health_Score": NUM,
          "Growth_Score": NUM, "Value_Score": NUM, "Sector_Rank": INT, "ROIC": PCT,
          "Gross_Margin": PCT, "FCF_Margin": PCT, "Revenue_CAGR_5y": PCT,
          "Earnings_Yield": PCT, "Net_Debt_EBITDA": RATIO}
ws = wb.create_sheet("Sector_Leaders")
style_header(ws, sl_cols, {"Company_Name": 26, "Why_It_Ranks": 40})
apply_num_format(ws, sl_cols, sl_fmt)

# 12-22. GICS sector sheets
sector_headers = core_names + [s[0] for s in SCORE_COLS]
sector_fmt = dict(core_fmt)
for s in SCORE_COLS:
    sector_fmt[s[0]] = s[1]
for sec in GICS_SECTORS:
    ws = wb.create_sheet(sec)
    style_header(ws, sector_headers)
    apply_num_format(ws, sector_headers, sector_fmt)
    # color scale on composite score (column index = len(core_names) + 6 = Composite_Score)
    comp_col = get_column_letter(len(core_names) + 6)
    color_scale(ws, comp_col)

wb.save(OUT)
print("Saved:", OUT)
print("Sheets:", wb.sheetnames)
print("Core columns:", len(core_names))
print("Sector metric definitions:", len(SECTOR_METRICS))
