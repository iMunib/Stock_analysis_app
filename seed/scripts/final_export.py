# -*- coding: utf-8 -*-
"""FINAL DATA-ONLY EXPORT (Company_ID overhaul).

- Adds permanent Company_ID = Country:Primary_Ticker:Exchange to Universe, Core_Financials,
  Time_Series, Source_Audit, and Data_Quality.
- Rebuilds Core_Financials + Time_Series from cached raw CSVs only (no 720 re-extraction).
- Total_Debt = ShortTermDebt + LongTermDebt fallback (Calculated); Equity = TotalEquity fallback.
- Preserves existing derived columns by Company_Name. No ranking/composite computation.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) + "/scripts")
sys.path.insert(0, "scripts")
import pandas as pd
import openpyxl

from common import XLSX, save_workbook_atomic
import assemble

RETRIEVAL = "2026-08-21"


def clean(v):
    return None if pd.isna(v) else v


# ---------------------------------------------------------------------------
# 1. Build
# ---------------------------------------------------------------------------
print("[1] building core + time series (Company_ID)...")
core = assemble.build_core()
ts = assemble.build_time_series()
univ = assemble.load_universe()
print(f"    core rows={len(core)} unique cid={core['Company_ID'].nunique()} | ts rows={len(ts)}")

# ---------------------------------------------------------------------------
# 2. Load workbook + existing Core for derived preservation
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(XLSX)
existing = pd.read_excel(XLSX, sheet_name="Core_Financials", keep_default_na=False)
existing_cols = list(existing.columns)
core_overlap = [c for c in assemble.CORE_COLS if c in existing_cols and c not in ("Company_ID",)]
new_cols = [c for c in assemble.CORE_COLS if c not in existing_cols]
derived_cols = [c for c in existing_cols if c not in assemble.CORE_COLS]

# name_map: Company_Name -> best (most-populated) derived values from existing sheet
name_map = {}
for _, r in existing.iterrows():
    cn = str(r["Company_Name"])
    key = "Rogers Communications Inc." if cn == "Rogers Communications" else cn
    cur = name_map.setdefault(key, {})
    cand = {c: r[c] for c in derived_cols if pd.notna(r[c])}
    if len(cand) > len(cur):
        name_map[key] = cand

# ---------------------------------------------------------------------------
# 3. Universe: insert Company_ID column (after Ticker, column index 3)
# ---------------------------------------------------------------------------
print("[3] adding Company_ID to Universe...")
ws_u = wb["Universe"]
u_hdr = [c.value for c in ws_u[1]]
if "Company_ID" not in u_hdr:
    ws_u.insert_cols(3)
    ws_u.cell(row=1, column=3, value="Company_ID")
    cid_col = [assemble.resolve_company_id(r["Country_of_Listing"], r["Ticker"], r.get("Exchange", ""))
               for _, r in univ.iterrows()]
    for i, v in enumerate(cid_col, start=2):
        ws_u.cell(row=i, column=3, value=v)
    print("    inserted at column 3")

# ---------------------------------------------------------------------------
# 4. Core_Financials full replace
# ---------------------------------------------------------------------------
print("[4] rebuilding Core_Financials...")
ws = wb["Core_Financials"]
final_hdr = ["Company_ID"] + existing_cols + new_cols
# overwrite header row
for i, h in enumerate(final_hdr, start=1):
    ws.cell(row=1, column=i, value=h)
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

written = 0
for _, row in core.iterrows():
    extra = name_map.get(str(row["Company_Name"]), {})
    vals = []
    for h in final_hdr:
        if h == "Company_ID":
            vals.append(row["Company_ID"])
        elif h in assemble.CORE_COLS and h in row.index:
            vals.append(clean(row[h]))
        elif h in extra:
            vals.append(clean(extra[h]))
        else:
            vals.append(None)
    ws.append(vals)
    written += 1
print(f"    rows written={written}")

# ---------------------------------------------------------------------------
# 5. Time_Series full replace (Company_ID first)
# ---------------------------------------------------------------------------
print("[5] rebuilding Time_Series...")
ts_ws = wb["Time_Series"]
ts_cols = ["Company_ID", "Ticker", "Fiscal_Year", "Field", "Value", "Unit", "Currency",
           "Source", "Retrieval_Date", "Notes"]
for i, h in enumerate(ts_cols, start=1):
    ts_ws.cell(row=1, column=i, value=h)
if ts_ws.max_row > 1:
    ts_ws.delete_rows(2, ts_ws.max_row - 1)
for row in ts.itertuples(index=False, name=None):
    ts_ws.append([clean(v) for v in row])
print(f"    rows written={len(ts)}")

# ---------------------------------------------------------------------------
# 6. Source_Audit: add Company_ID column (map from Ticker)
# ---------------------------------------------------------------------------
print("[6] adding Company_ID to Source_Audit...")
# build ticker -> Company_ID map (bare + yahoo)
tk2cid = {}
for _, u in univ.iterrows():
    cid = assemble.resolve_company_id(u["Country_of_Listing"], u["Ticker"], u.get("Exchange", ""))
    tk2cid[str(u["Ticker"])] = cid
    tk2cid[str(u["Yahoo_Ticker"])] = cid
    tk2cid[str(u["Primary_Ticker"])] = cid
sa_ws = wb["Source_Audit"]
sa_hdr = [c.value for c in sa_ws[1]]
if "Company_ID" not in sa_hdr:
    sa_ws.insert_cols(2)
    sa_ws.cell(row=1, column=2, value="Company_ID")
    for r in range(2, sa_ws.max_row + 1):
        tk = sa_ws.cell(row=r, column=1).value
        sa_ws.cell(row=r, column=2, value=tk2cid.get(str(tk), ""))
    print("    inserted at column 2")

# ---------------------------------------------------------------------------
# 7. Data_Quality: populate N/A reasons for non-COMPLETE rows
# ---------------------------------------------------------------------------
print("[7] populating Data_Quality...")
dq_ws = wb["Data_Quality"]
dq_hdr = ["Company_ID", "Ticker", "Field", "Issue", "Source_Attempted", "Retrieval_Date", "Resolution", "Notes"]
for i, h in enumerate(dq_hdr, start=1):
    dq_ws.cell(row=1, column=i, value=h)
if dq_ws.max_row > 1:
    dq_ws.delete_rows(2, dq_ws.max_row - 1)

FIVE = ["Revenue", "Net_Income", "Book_Equity", "Total_Debt", "Operating_Cash_Flow"]
dq_rows = 0
for _, row in core.iterrows():
    st = row["Extraction_Status"]
    if st == "COMPLETE":
        continue
    if st == "MISSING_SOURCE":
        dq_ws.append([row["Company_ID"], row["Ticker"], "ALL", "MISSING_SOURCE",
                      "EDGAR/Yahoo", RETRIEVAL,
                      "No source financials available (CIK/ticker extraction failed or no filing yet)",
                      ""])
        dq_rows += 1
    else:  # PARTIAL
        missing = [f for f in FIVE if row.get(f) is None or (isinstance(row.get(f), str) and row.get(f).strip() == "")]
        for f in missing:
            dq_ws.append([row["Company_ID"], row["Ticker"], f, "MISSING_FIELD",
                          "EDGAR/Yahoo", RETRIEVAL, "Field tag not present in source; kept N/A", ""])
            dq_rows += 1
print(f"    data_quality rows={dq_rows}")

# ---------------------------------------------------------------------------
# 8. Save atomically
# ---------------------------------------------------------------------------
print("[8] saving...")
save_workbook_atomic(wb)
print("    saved.")
