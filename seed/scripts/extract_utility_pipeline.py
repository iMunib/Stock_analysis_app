# -*- coding: utf-8 -*-
"""Utility/pipeline sheet writer.
Reads raw/utility_pipeline_source.csv (recapturable source) and writes the
Utility_Pipeline sheet. Issuer cash-metric terms and non-GAAP values preserved verbatim;
never invents FFO/rate-base/DCF.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill

from common import RAW, XLSX

SRC = os.path.join(RAW, "utility_pipeline_source.csv")

UP_COLS = ["Ticker", "Company_Name", "Peer_Group", "Fiscal_Period_End", "FFO_or_Equivalent",
           "Issuer_Term", "FFO_to_Debt", "Debt_to_EBITDA", "Interest_Coverage", "Dividend_Payout",
           "Capex", "Rate_Base", "Rate_Base_Growth", "Allowed_ROE", "DCF_or_Comparable",
           "Coverage_Ratio", "Commodity_vs_Contracted_Mix", "Source_URL", "Confidence", "Notes"]


def write():
    df = pd.read_csv(SRC, keep_default_na=False)
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Utility_Pipeline"]
    hfill = PatternFill("solid", fgColor="1F4E78"); hfont = Font(bold=True, color="FFFFFF")
    for j, h in enumerate(UP_COLS, start=1):
        c = ws.cell(row=1, column=j, value=h); c.fill = hfill; c.font = hfont
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for _, row in df.iterrows():
        for j, h in enumerate(UP_COLS, start=1):
            v = row.get(h, "")
            ws.cell(row=r, column=j, value=v if v != "" else None)
        r += 1
    wb.save(XLSX)
    return len(df)


if __name__ == "__main__":
    print("Utility_Pipeline rows written:", write())
