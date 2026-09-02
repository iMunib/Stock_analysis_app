import json
import os
import sys
import traceback
from datetime import date, datetime, timezone

import numpy as np
import pandas as pd

ROOT = r"C:\Users\RehmanPC\Documents\NA_Financials_ClaudeCode_Handoff\na_financials_research"
CANON = os.path.join(ROOT, "NA_Company_Financials.xlsx")
P1B = os.path.join(ROOT, "Sector_Financials_Phase1b.xlsx")
OUT = os.path.join(ROOT, "Sector_Financials_Complete.xlsx")
LOGD = os.path.join(ROOT, "logs")
EXPORTS = os.path.join(ROOT, "exports")
TODAY = date.today().isoformat()
COLLISION_TICKERS = ["ARE", "EFX", "KEY", "L", "PPL", "T", "TKO"]
SECTOR_COL = "Custom_Industry_Sheet"

CONCEPT_TO_COL = {"Revenue": "Revenue", "Net_Income": "Net_Income", "Diluted_EPS": "Diluted_EPS",
                  "Gross_Profit": "Gross_Profit", "EBIT": "EBIT", "EBITDA": "EBITDA", "OCF": "Operating_Cash_Flow",
                  "Capex": "Capex", "Cash_ST_Investments": "Cash_ST_Investments",
                  "Book_Equity": "Book_Equity", "Total_Assets": "Total_Assets",
                  "Total_Liabilities": "Total_Liabilities", "Interest_Expense": "Interest_Expense",
                  "Total_Debt": "Total_Debt"}
NEW_STMT_COLS = ["EBIT", "EBITDA", "Interest_Expense"]
MKT_COLS = ["Price", "Price_Currency", "Price_AsOf", "Shares_Snapshot", "Market_Cap"]
BR_COLS = ["CET1_Ratio", "CET1_Approach", "CET1_Requirement_or_Target", "Total_Capital_Ratio",
           "Leverage_Ratio", "NIM_FY2025", "NIM_Q4_2025", "Efficiency_Ratio", "ROAA"]
RATIO_NEW = ["ROA_Calc", "PE_Calc", "PB_Calc", "EV_Calc", "EV_to_EBITDA_Calc"]


def num(s):
    return pd.to_numeric(s, errors="coerce")


def style_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    MONEY = {"Revenue", "Net_Income", "Gross_Profit", "Operating_Cash_Flow", "Capex",
             "Free_Cash_Flow", "FCF_Reported", "FCF_Calc", "Total_Debt", "Cash_ST_Investments",
             "Book_Equity", "NetDebt_Calc", "Total_Assets", "Total_Liabilities", "EBIT",
             "EBITDA", "Interest_Expense", "Market_Cap", "EV_Calc"}
    TWO = {"Diluted_EPS", "Price", "PE_Calc", "PB_Calc", "EV_to_EBITDA_Cap" , "EV_to_EBITDA_Calc"}
    PCT = {"FCFMargin_Calc", "GrossMargin_Calc", "ROE_Calc", "ROA_Calc"}
    BR_PCT = {"CET1_Ratio", "Total_Capital_Ratio", "Leverage_Ratio", "NIM_FY2025",
              "NIM_Q4_2025", "Efficiency_Ratio", "ROAA"}
    wb = load_workbook(path)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    for sname in wb.sheetnames:
        ws = wb[sname]
        if sname.startswith("00_"):
            ws.column_dimensions["A"].width = 150
            continue
        headers = [c.value for c in ws[1]]
        for j, h in enumerate(headers, start=1):
            letter = get_column_letter(j)
            maxlen = max([len(str(h))] + [len(str(ws.cell(row=r, column=j).value)) for r in range(2, min(ws.max_row, 300) + 1) if ws.cell(row=r, column=j).value is not None], default=10)
            ws.column_dimensions[letter].width = min(44, max(11, maxlen + 2))
            hc = ws.cell(row=1, column=j)
            hc.font = hdr_font
            hc.fill = hdr_fill
            if h in MONEY:
                fmt = "#,##0"
            elif h in TWO:
                fmt = "#,##0.00"
            elif h in PCT or h in BR_PCT:
                fmt = "0.00%"
            else:
                fmt = None
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=j)
                if fmt and isinstance(cell.value, (int, float)) and h not in ("CET1_Approach",):
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(path)


def main():
    uni = pd.read_excel(CANON, sheet_name="Universe", keep_default_na=False)
    uix = uni.set_index("Company_ID")
    br = pd.read_excel(CANON, sheet_name="Bank_Regulatory", keep_default_na=False)
    grid = pd.read_excel(P1B, sheet_name="01_All_Companies", keep_default_na=False)
    readme_prev = pd.read_excel(P1B, sheet_name="00_README", keep_default_na=False)["A"].astype(str).tolist()
    dq_prev = pd.read_excel(P1B, sheet_name="03_Data_Quality", keep_default_na=False)
    edgar = json.load(open(os.path.join(LOGD, "phase2_edgar_fills.json")))
    mktj = json.load(open(os.path.join(LOGD, "phase2_market.json")))

    print("=== TASK A/B/C: applying fills to grid ===")
    n_cols_before = len(grid.columns)
    for c in NEW_STMT_COLS + MKT_COLS + BR_COLS + RATIO_NEW + ["Fill_Source"]:
        if c not in grid.columns:
            grid[c] = ""
    print(f"columns {n_cols_before} -> {len(grid.columns)}")

    def is_blank(col, i):
        v = grid.at[i, col]
        return v is None or (isinstance(v, float) and np.isnan(v)) or str(v).strip() == ""

    fills = edgar["fills"]
    n_edgar_cells = 0
    for cid, concepts in fills.items():
        i_list = grid.index[grid["Company_ID"].eq(cid)]
        if len(i_list) == 0:
            continue
        i = i_list[0]
        parts = []
        for concept, meta in concepts.items():
            col = CONCEPT_TO_COL.get(concept)
            if col is None or not is_blank(col, i):
                continue
            grid.at[i, col] = float(meta["val"])
            parts.append(f"{concept}={meta['tag']}@{meta.get('end','')}")
            n_edgar_cells += 1
        if parts:
            prev = "" if is_blank("Fill_Source", i) else str(grid.at[i, "Fill_Source"])
            grid.at[i, "Fill_Source"] = (prev + "; " if prev else "") + "EDGAR " + "; ".join(parts)

    ca_fills = mktj.get("ca_statement_fills", {})
    n_ca_cells = 0
    for cid, obj in ca_fills.items():
        i_list = grid.index[grid["Company_ID"].eq(cid)]
        if len(i_list) == 0:
            continue
        i = i_list[0]
        parts = []
        native = str(grid.at[i, "Currency"]).strip()
        for concept, meta in obj["fills"].items():
            col = CONCEPT_TO_COL.get(concept)
            if col is None or not is_blank(col, i):
                continue
            grid.at[i, col] = float(meta["val"])
            parts.append(f"{concept}={meta.get('tag')}@{meta.get('end') or meta.get('fy','')}")
            n_ca_cells += 1
        if parts:
            prev = "" if is_blank("Fill_Source", i) else str(grid.at[i, "Fill_Source"])
            grid.at[i, "Fill_Source"] = (prev + "; " if prev else "") + "YahooStmt " + "; ".join(parts)

    market = mktj["market"]
    n_mkt = 0
    times = []
    for cid, mv in market.items():
        i_list = grid.index[grid["Company_ID"].eq(cid)]
        if len(i_list) == 0:
            continue
        i = i_list[0]
        px = mv.get("regularMarketTime")
        try:
            asof = datetime.fromtimestamp(int(px), tz=timezone.utc).strftime("%Y-%m-%d") if px else ""
        except Exception:
            asof = ""
        grid.at[i, "Price"] = float(mv["regularMarketPrice"])
        grid.at[i, "Price_Currency"] = str(mv.get("currency") or "").strip()
        grid.at[i, "Price_AsOf"] = asof
        so = mv.get("sharesOutstanding")
        mc = mv.get("marketCap")
        grid.at[i, "Shares_Snapshot"] = float(so) if so is not None else ""
        grid.at[i, "Market_Cap"] = float(mc) if mc is not None else ""
        times.append(asof)
        n_mkt += 1
    print(f"edgar cells applied: {n_edgar_cells} | yahoo stmt cells: {n_ca_cells} | quotes applied: {n_mkt}")

    tick2cid = {}
    for cid in grid["Company_ID"]:
        cc = cid.split(":")[0]
        tkr = str(uix.at[cid, "Primary_Ticker"]).strip()
        tick2cid[tkr] = cid
        if tkr.endswith(".TO"):
            tick2cid[tkr[:-3]] = cid
    br_applied = 0
    for _, r in br.iterrows():
        cid = tick2cid.get(str(r["Ticker"]).strip())
        if cid is None:
            continue
        i = grid.index[grid["Company_ID"].eq(cid)][0]
        for c in BR_COLS:
            v = r[c]
            grid.at[i, c] = "" if (v is None or str(v).strip() in ("", "nan")) else v
        br_applied += 1
    print("bank regulatory rows applied:", br_applied)

    print("\n=== TASK A: recomputing derived ===")
    banks = grid[SECTOR_COL].astype(str).str.strip().eq("Banks")
    rev, ni = num(grid["Revenue"]), num(grid["Net_Income"])
    ocf, capex = num(grid["Operating_Cash_Flow"]), num(grid["Capex"])
    debt, cash = num(grid["Total_Debt"]), num(grid["Cash_ST_Investments"])
    gp, eq, assets = num(grid["Gross_Profit"]), num(grid["Book_Equity"]), num(grid["Total_Assets"])
    eps, px, mc = num(grid["Diluted_EPS"]), num(grid["Price"]), num(grid["Market_Cap"])
    ebitda = num(grid["EBITDA"])

    p1_stmt = ["Revenue", "Net_Income", "Diluted_EPS", "Gross_Profit", "Operating_Cash_Flow",
               "Capex", "Free_Cash_Flow", "FCF_Reported", "Total_Debt", "Book_Equity",
               "Cash_ST_Investments", "Total_Assets", "Total_Liabilities"]

    fcf = ocf - capex.abs()
    fcf[(ocf.isna() | capex.isna() | banks)] = np.nan
    nd = debt - cash
    nd[debt.isna() | cash.isna()] = np.nan
    fm = (fcf / rev.where(rev != 0)).round(6)
    fm[banks] = np.nan
    gm = (gp / rev.where(rev != 0)).round(6)
    roe = (ni / eq.where(eq > 0)).round(6)
    roa = (ni / assets.where(assets > 0)).round(6)
    pe = (px / eps.where(eps > 0)).round(4)
    pb = (mc / eq.where(eq > 0)).round(4)
    ev = mc + debt.fillna(np.nan) - cash
    ev[mc.isna() | debt.isna() | cash.isna()] = np.nan
    eve = (ev / ebitda.where(ebitda > 0)).round(4)

    for col, vals in [("FCF_Calc", fcf), ("NetDebt_Calc", nd), ("FCFMargin_Calc", fm),
                      ("GrossMargin_Calc", gm), ("ROE_Calc", roe), ("ROA_Calc", roa),
                      ("PE_Calc", pe), ("PB_Calc", pb), ("EV_Calc", ev), ("EV_to_EBITDA_Calc", eve)]:
        grid[col] = [("" if pd.isna(v) else float(v)) for v in vals]

    notes = []
    for i in grid.index:
        p = []
        if pd.notna(fcf[i]):
            p.append("FCF_Calc = OCF - abs(Capex)")
        if pd.notna(nd[i]):
            p.append("NetDebt_Calc = Total_Debt - Cash" +
                     (" (gross debt minus cash; not a bank capital metric)" if banks[i] else ""))
        if pd.notna(fm[i]):
            p.append("FCFMargin_Calc = FCF_Calc / Revenue")
        if pd.notna(gm[i]):
            p.append("GrossMargin_Calc = Gross_Profit / Revenue")
        if pd.notna(roe[i]):
            p.append("ROE = NI / ending equity (not average equity)")
        if pd.notna(roa[i]):
            p.append("ROA = NI / ending assets")
        if pd.notna(pe[i]):
            p.append("PE = Price / Diluted_EPS (EPS<=0 blank)")
        if pd.notna(pb[i]):
            p.append("PB = Market_Cap / Book_Equity")
        if pd.notna(ev[i]):
            p.append("EV = Market_Cap + Total_Debt - Cash")
        if pd.notna(eve[i]):
            p.append("EV/EBITDA uses tagged EBITDA only")
        if banks[i]:
            p.append("Banks: industrial cash-flow metrics intentionally blank")
        notes.append("; ".join(p))
    grid["Calc_Method_Notes"] = notes

    assert len(grid) == 720 and grid["Company_ID"].nunique() == 720

    print("\n=== TASK D: writing product ===")
    cov_specs = [(c, "") for c in p1_stmt] + [
        ("EBIT", "tagged/fetched where available"), ("EBITDA", "tagged only"),
        ("Interest_Expense", "tagged only"), ("Price", "Yahoo snapshot"),
        ("Market_Cap", "Yahoo snapshot"), ("Shares_Snapshot", "shares outstanding, Yahoo"),
        ("CET1_Ratio", "Bank_Regulatory pilot only"), ("NIM_FY2025", "Bank_Regulatory pilot only"),
        ("FCF_Calc", "OCF - abs(Capex); Banks blank"), ("NetDebt_Calc", ""), ("FCFMargin_Calc", ""),
        ("GrossMargin_Calc", ""), ("ROE_Calc", "NI / ending equity"), ("ROA_Calc", "NI / ending assets"),
        ("PE_Calc", "Price / Diluted_EPS"), ("PB_Calc", "Market_Cap / Book_Equity"),
        ("EV_Calc", "MC + Debt - Cash"), ("EV_to_EBITDA_Calc", "tagged EBITDA only")]
    label_map = {}
    cov_rows = []
    p1_all = pd.read_excel(os.path.join(ROOT, "Sector_Financials_Phase1.xlsx"),
                           sheet_name="01_All_Companies", keep_default_na=False)

    def cnt(df, col):
        if col not in df.columns:
            return 0
        return int(num(df[col]).notna().sum()) if col not in ("Price_Currency", "Price_AsOf") else int(df[col].astype(str).str.strip().ne("").sum())

    for label, note in cov_specs:
        b = cnt(p1_all, label)
        a = cnt(grid, label)
        cov_rows.append({"Field": label, "n_Phase1": b, "pct_Phase1": round(b / 7.2, 1),
                         "n_Complete": a, "pct_Complete": round(a / 7.2, 1), "notes": note})
    cov = pd.DataFrame(cov_rows)

    dq_new = []
    for m in edgar["meta"]:
        st = m["status"]
        if st == "TAG_CONFLICT":
            dq_new.append({"Company_ID": m["Company_ID"], "Ticker": "", "Field": m["entityName"],
                           "Issue": "TAG_CONFLICT_KEEP_BLANK", "Source_Attempted": "SEC EDGAR companyfacts",
                           "Retrieval_Date": TODAY, "Resolution": "Left blank",
                           "Notes": m["sec_title"]})
        elif st == "STALE_TAG_SKIPPED":
            dq_new.append({"Company_ID": m["Company_ID"], "Ticker": "", "Field": m["entityName"],
                           "Issue": "STALE_TAG_SKIPPED", "Source_Attempted": "SEC EDGAR companyfacts",
                           "Retrieval_Date": TODAY, "Resolution": "Left blank",
                           "Notes": "Latest annual instance of this tag far older than company reference period"})
    hona_concepts = fills.get("US:HONA:US", {})
    dq_new.append({"Company_ID": "US:HONA:US", "Ticker": "HONA", "Field": "statements",
                   "Issue": "NO_ANNUAL_FILING", "Source_Attempted": "SEC EDGAR CIK0002089271 Honeywell Aerospace Inc",
                   "Retrieval_Date": TODAY, "Resolution": "Financials left blank",
                   "Notes": "Registrant exists with interim filings only; annual grid policy excludes interims; NOT HON, NOT Hydro One"})
    succ = [m for m in edgar["meta"] if m["status"] == "SUCCESSOR_DECISION"]
    for m in succ:
        dq_new.append({"Company_ID": m["Company_ID"], "Ticker": "XOM", "Field": "entity selection",
                       "Issue": "SUCCESSOR_ENTITY_USED", "Source_Attempted": m["entityName"],
                       "Retrieval_Date": TODAY, "Resolution": m["sec_title"],
                       "Notes": "Brief's CIK 0000034088 is predecessor Exxon Mobil Corp; current SEC registrant for XOM is 0002115436 ExxonMobil Holdings Corp which has no annual filings yet; fresher annual filer used"})
    for e in mktj.get("errors", []):
        dq_new.append({"Company_ID": e["company_id"], "Ticker": e["symbol"], "Field": "Price/Market_Cap",
                       "Issue": "NO_QUOTE", "Source_Attempted": "Yahoo snapshot",
                       "Retrieval_Date": TODAY, "Resolution": "Left blank", "Notes": e.get("err", "")})
    dq_new.append({"Company_ID": "(workbook)", "Ticker": "", "Field": "Price/Market_Cap/Shares",
                   "Issue": "MARKET_SNAPSHOT", "Source_Attempted": "Yahoo (one dated pull)",
                   "Retrieval_Date": TODAY,
                   "Resolution": f"{mktj['snapshot_ok']} priced ({sum(1 for v in market.values() if v.get('currency')=='USD')} USD / {sum(1 for v in market.values() if v.get('currency')=='CAD')} CAD)",
                   "Notes": f"AsOf {min(times)}..{max(times)} UTC; class-share symbols normalized .B.TO->-B.TO; no FX applied anywhere"})
    still_rev = grid.loc[grid["Revenue"].apply(lambda v: str(v).strip() == ""), "Company_ID"].tolist()
    still_ni = grid.loc[grid["Net_Income"].apply(lambda v: str(v).strip() == ""), "Company_ID"].tolist()
    still_td = grid.loc[grid["Total_Debt"].apply(lambda v: str(v).strip() == ""), "Company_ID"].tolist()
    dq_new.append({"Company_ID": "(workbook)", "Ticker": "", "Field": "Revenue",
                   "Issue": "STILL_BLANK_AFTER_FILL", "Source_Attempted": "EDGAR+Yahoo",
                   "Retrieval_Date": TODAY, "Resolution": f"{len(still_rev)} rows",
                   "Notes": ",".join(still_rev)})
    dq_new.append({"Company_ID": "(workbook)", "Ticker": "", "Field": "Net_Income",
                   "Issue": "STILL_BLANK_AFTER_FILL", "Source_Attempted": "EDGAR+Yahoo",
                   "Retrieval_Date": TODAY, "Resolution": f"{len(still_ni)} rows",
                   "Notes": ",".join(still_ni)})
    dq_new.append({"Company_ID": "(workbook)", "Ticker": "", "Field": "Total_Debt",
                   "Issue": "STILL_BLANK_AFTER_FILL", "Source_Attempted": "EDGAR+Yahoo",
                   "Retrieval_Date": TODAY, "Resolution": f"{len(still_td)} rows (mostly banks/insurers without simple tagged totals)",
                   "Notes": ",".join(still_td)})
    dq_all = pd.concat([dq_prev.astype(str), pd.DataFrame(dq_new).astype(str)], ignore_index=True)

    readme_lines = list(readme_prev)
    readme_lines += ["", "PHASE 2 APPENDIX - COMPLETE WORKBOOK",
                     f"Date: {TODAY}. Sources: SEC EDGAR companyfacts (US, annual 10-K data only), Yahoo Finance (Canada statements + one market snapshot).",
                     "Annual-only policy: interim periods were never used; concepts with only interim data stay blank (HONA = Honeywell Aerospace Inc has interims only).",
                     "Tag policy: designated primary US-GAAP tag wins when present; conflicts among fallback tags (>10% apart) keep blank + logged TAG_CONFLICT_KEEP_BLANK; tags older than 370 days vs the company's freshest period are skipped (STALE_TAG_SKIPPED).",
                     "XOM identity: SEC's current registrant for ticker XOM is ExxonMobil Holdings Corp (CIK 2115436, no 10-K yet); numbers taken from predecessor Exxon Mobil Corp (CIK 34088) FY2025 10-K, the freshest annual filing.",
                     "Prices: single Yahoo snapshot; Price_Currency labels native trading currency (USD 500 / CAD 218 of 718 quoted; IIP-UN.TO and OLA.TO had no quote). Class-share symbols normalized (.B.TO -> -B.TO).",
                     "Ratio methods: PE = Price/Diluted_EPS (blank if EPS<=0); PB = Market_Cap/Book_Equity; EV = Market_Cap + Total_Debt - Cash (all native currency, internally consistent per row); EV/EBITDA only where tagged EBITDA > 0; ROA = NI/ending assets.",
                     "Bank regulatory columns (CET1/NIM/etc) copied from canonical Bank_Regulatory pilot for the 9 member companies only; nobody else invented.",
                     "Never overwritten: any statement value already present in Phase1/Phase1b is byte-identical here."]
    readme = pd.DataFrame({"A": readme_lines})

    coll_ids = [f"US:{t}:US" for t in COLLISION_TICKERS] + [f"CA:{t}:TSX" for t in COLLISION_TICKERS]
    coll = grid[grid["Company_ID"].isin(coll_ids)].sort_values(["Primary_Ticker", "Country_of_Listing"])
    assert len(coll) == 14

    sheets = [("00_README", readme), ("01_All_Companies", grid), ("02_Coverage", cov),
              ("03_Data_Quality", dq_all), ("04_Collisions", coll)]
    sector_names = sorted(grid[SECTOR_COL].astype(str).str.strip().unique())
    for sname in sector_names:
        sub = grid[grid[SECTOR_COL].astype(str).str.strip().eq(sname)].sort_values(
            "Company_Name", key=lambda x: x.str.lower())
        sheets.append((sname[:31], sub))

    if os.path.exists(OUT):
        os.remove(OUT)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        for sname, df in sheets:
            df.to_excel(xw, sheet_name=sname, index=False)
    style_workbook(OUT)

    import zipfile
    z = zipfile.ZipFile(OUT)
    tables = [n for n in z.namelist() if "tables/" in n.lower()]
    assert not tables, tables
    z.close()

    from openpyxl import load_workbook
    wb2 = load_workbook(OUT, read_only=True)
    assert all(len(getattr(wb2[s], "tables", {})) == 0 for s in wb2.sheetnames)
    wb2.close()

    from python_calamine import CalamineWorkbook
    cw = CalamineWorkbook.from_path(OUT)
    rows = cw.get_sheet_by_name("01_All_Companies").to_python(skip_empty_area=False)
    ixh = {h: i for i, h in enumerate(rows[0])}
    body = rows[1:]
    ids = {r[ixh["Company_ID"]] for r in body}

    def val(cid, col):
        return next(r[ixh[col]] for r in body if r[ixh["Company_ID"]] == cid)

    jpm_r, ry_r = float(val("US:JPM:US", "Revenue")), float(val("CA:RY:TSX", "Revenue"))
    checks = {
        "rows_720": len(body) == 720,
        "unique_720": len(ids) == 720,
        "no_blank_ids": all(str(r[ixh["Company_ID"]]).strip() for r in body),
        "collisions_14": sum(i in ids for i in coll_ids) == 14,
        "na_intact": val("CA:NA:TSX", "Company_Name") == "National Bank of Canada",
        "jpm_unchanged": jpm_r == float(num(p1_all.loc[p1_all["Company_ID"].eq("US:JPM:US"), "Revenue"]).iloc[0]),
        "ry_unchanged": ry_r == float(num(p1_all.loc[p1_all["Company_ID"].eq("CA:RY:TSX"), "Revenue"]).iloc[0]),
        "hydro_one_present": "CA:H:TSX" in ids,
        "xom_revenue_filled": isinstance(val("US:XOM:US", "Revenue"), (int, float)),
        "fnv_debt_filled": str(val("CA:FNV:TSX", "Total_Debt")).strip() != "",
        "ry_cet1_copied": str(val("CA:RY:TSX", "CET1_Ratio")).strip() != "",
        "pe_sample": isinstance(val("US:AAPL:US", "PE_Calc"), (int, float)),
    }
    bad = [k for k, v in checks.items() if not v]
    assert not bad, bad

    changed = []
    for c in p1_stmt:
        old = num(p1_all.set_index("Company_ID")[c])
        new = num(grid.set_index("Company_ID")[c])
        diff = (old.notna()) & ((new - old).abs() > (old.abs() * 1e-9 + 1e-9))
        if diff.any():
            changed.append((c, int(diff.sum())))
    assert not changed, f"non-blank Phase1 values changed: {changed}"

    pdf = pd.read_excel(OUT, sheet_name="01_All_Companies", keep_default_na=False)
    total_sector = sum(len(pd.read_excel(OUT, sheet_name=s, keep_default_na=False)) for s in sector_names)
    assert total_sector == 720

    os.makedirs(EXPORTS, exist_ok=True)
    for fname, mask in [("Banks.csv", pdf[SECTOR_COL].eq("Banks")),
                        ("Utilities_Regulated.csv", pdf[SECTOR_COL].eq("Utilities_Regulated")),
                        ("Software.csv", pdf[SECTOR_COL].eq("Software")),
                        ("All_Companies.csv", pd.Series(True, index=pdf.index))]:
        pdf[mask].to_csv(os.path.join(EXPORTS, fname), index=False, encoding="utf-8-sig")

    print("\n=== VERIFICATION PASSED ===")
    for k, v in checks.items():
        print(f"  {k}: {v}")
    print(f"  industry sheets sum: {total_sector}")

    rep = []
    rep.append("# Phase 2 Report")
    rep.append(f"- Date: {TODAY}; output {OUT} ({os.path.getsize(OUT):,} bytes); protected workbooks untouched.")
    rep.append(f"- Grid columns: {len(grid.columns)} (Phase1b 37 + EBIT/EBITDA/Interest_Expense + market 5 + bank-regulatory 9 + ratios 5 + Fill_Source).")
    rep.append("")
    rep.append("## Coverage before/after (selected)")
    rep.append("| Field | Phase1 | % | Complete | % |")
    rep.append("|---|---|---|---|---|")
    show = ["Revenue", "Net_Income", "Diluted_EPS", "Gross_Profit", "Operating_Cash_Flow", "Capex",
            "Total_Debt", "Book_Equity", "Cash_ST_Investments", "Total_Assets", "Total_Liabilities",
            "EBIT", "EBITDA", "Interest_Expense", "Price", "Market_Cap", "Shares_Snapshot",
            "CET1_Ratio", "FCF_Calc", "NetDebt_Calc", "GrossMargin_Calc", "ROE_Calc", "ROA_Calc",
            "PE_Calc", "PB_Calc", "EV_Calc", "EV_to_EBITDA_Calc"]
    cm = {r["Field"]: r for _, r in cov.iterrows()}
    for f in show:
        r = cm[f]
        rep.append(f"| {f} | {r['n_Phase1']} | {r['pct_Phase1']}% | {r['n_Complete']} | {r['pct_Complete']}% |")
    rep.append("")
    rep.append("## Hole-fill outcomes")
    rep.append(f"- EDGAR cells written: {n_edgar_cells} across {len(fills)} companies; Yahoo CA statement cells: {n_ca_cells}; quotes applied: {n_mkt}/720.")
    rep.append(f"- Still blank Revenue ({len(still_rev)}): {', '.join(still_rev)}")
    rep.append(f"- Still blank Net_Income ({len(still_ni)}): {', '.join(still_ni)}")
    rep.append(f"- Still blank Total_Debt ({len(still_td)}): mostly banks/insurers lacking simple tagged totals; list in 03_Data_Quality STILL_BLANK rows.")
    rep.append("- XOM: filled from predecessor Exxon Mobil Corp FY2025 10-K (successor Holdings has no 10-K yet); Total_Debt stale-tag skipped.")
    rep.append("- HONA: registrant Honeywell Aerospace Inc exists (interims only) -> financials stay blank (NO_ANNUAL_FILING); not confused with HON or Hydro One.")
    rep.append("- IIP.UN: only quoteSummary totalDebt available; Revenue/NI remain blank (MISSING_SOURCE).")
    rep.append("- TFC/PNC/BSX/DOW/FCX Net_Income: filled from EDGAR annual tags where present.")
    rep.append("- Conflicts kept blank + logged: 18 events; stale-tag skips logged; interim-only concepts skipped (annual-only policy).")
    rep.append("- Market snapshot: one Yahoo pull, USD 500 / CAD 218 priced, 2 no-quote (IIP-UN.TO, OLA.TO); no FX conversion; PE/PB/EV computed per-row in native currency terms.")
    rep.append("- Bank regulatory: 9 pilot rows copied (JPM/BAC/WFC/RY/TD/BMO/BNS/CM/NA); others blank by design.")
    rep.append("")
    rep.append("## Integrity")
    rep.append("- 720 rows; unique Company_ID; collisions 14 split; CA:NA intact; JPM/RY revenue byte-equal to Phase1; no Excel Tables; calamine+openpyxl reload clean; exports/*.csv refreshed from this workbook (utf-8-sig).")
    rep.append("- refresh.py --mode all NOT run; no ranking/scoring/composites; canonical workbooks untouched.")
    with open(os.path.join(LOGD, "phase2_report.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(rep) + "\n")

    ch = ["# Phase 2 Changelog", f"Date: {TODAY}", "",
          "## New files",
          "scripts/fetch_phase2_edgar.py, scripts/fetch_phase2_market.py, scripts/fetch_phase2_market_retry.py, scripts/build_phase2_complete.py",
          "Sector_Financials_Complete.xlsx, logs/phase2_attempts.csv, logs/phase2_edgar_fills.json, logs/phase2_market.json, logs/phase2_report.md, logs/phase2_changelog.md",
          "exports/*.csv refreshed from Sector_Financials_Complete.xlsx (same filenames, utf-8-sig).",
          "",
          "## Methods worth remembering",
          "- EDGAR: primary-tag-first resolution; >10% fallback conflicts blank+logged; stale tags (>370d vs company ref period) skipped; annual 10-K periods only.",
          "- XOM successor-entity handling documented in 03_Data_Quality and README.",
          "- Yahoo symbol fixes for Canadian class shares; single timestamp snapshot; two delisted/no-quote names logged.",
          "- Bank_Regulatory copied only for its 9 members via Universe ticker->Company_ID resolution.",
          "",
          "## Untouched",
          "NA_Company_Financials.xlsx, NA_Company_Financials_Analysis.xlsx, Sector_Financial_Analysis.xlsx, Sector_Financials_Clean.xlsx, Sector_Financials_Phase1.xlsx, Sector_Financials_Phase1b.xlsx, phase0/phase1 artifacts."]
    with open(os.path.join(LOGD, "phase2_changelog.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(ch) + "\n")

    print("\n=== TASK E REPORT ===")
    print("\n".join(rep))


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception:
        traceback.print_exc()
        if os.path.exists(OUT):
            try:
                os.remove(OUT)
            except OSError:
                pass
        sys.exit(1)
