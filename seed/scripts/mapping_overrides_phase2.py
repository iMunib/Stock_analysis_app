# -*- coding: utf-8 -*-
"""
Phase 2 prep — manual mapping override table + apply overrides to Universe.
"""
import pandas as pd, os
import openpyxl

ROOT = r"C:\Users\RehmanPC\.openclaw-autoclaw\workspace\na_financials_research"
XLSX = os.path.join(ROOT, "NA_Company_Financials.xlsx")

# (ticker, country, new Custom_Industry_Sheet)
OVERRIDES = [
    ("ABBV", "US", "Pharma"),
    ("AC",   "CA", "Airlines"),
    ("ACN",  "US", "Software"),
    ("ENB",  "CA", "Pipelines_Midstream"),
    ("TRP",  "CA", "Pipelines_Midstream"),
    ("CNR",  "CA", "Railroads"),
    ("CP",   "CA", "Railroads"),
    ("BCE",  "CA", "Telecom"),
    ("T",    "CA", "Telecom"),
    ("RCI.B","CA", "Telecom"),
    ("H",    "CA", "Utilities_Regulated"),
    ("FTS",  "CA", "Utilities_Regulated"),
    ("EMA",  "CA", "Utilities_Regulated"),
]
REASON = "User Phase-2 directive"
SOURCE = "User Phase-2 requirement"

# load universe
wb = openpyxl.load_workbook(XLSX)
ws = wb["Universe"]
hdr = [c.value for c in ws[1]]
idx = {h: i+1 for i, h in enumerate(hdr)}

rows = []
applied = []
for r in range(2, ws.max_row + 1):
    ticker = ws.cell(row=r, column=idx["Ticker"]).value
    country = ws.cell(row=r, column=idx["Country_of_Listing"]).value
    for ot, oc, new_sheet in OVERRIDES:
        if ticker == ot and country == oc:
            old = ws.cell(row=r, column=idx["Custom_Industry_Sheet"]).value
            sector = ws.cell(row=r, column=idx["GICS_Sector"]).value
            sub = ws.cell(row=r, column=idx["GICS_Sub_Industry"]).value or ws.cell(row=r, column=idx["GICS_Industry"]).value
            name = ws.cell(row=r, column=idx["Company_Name"]).value
            rows.append([ot, name, sector, sub, new_sheet, REASON, SOURCE, "Applied"])
            ws.cell(row=r, column=idx["Custom_Industry_Sheet"]).value = new_sheet
            applied.append((ot, old, new_sheet))
            break

# write Mapping_Overrides sheet
if "Mapping_Overrides" in wb.sheetnames:
    del wb["Mapping_Overrides"]
mo = wb.create_sheet("Mapping_Overrides")
mo_cols = ["Ticker", "Company_Name", "GICS_Sector", "GICS_Sub_Industry",
           "Custom_Industry_Sheet", "Mapping_Reason", "Mapping_Source", "Review_Status"]
mo.append(mo_cols)
for row in rows:
    mo.append(row)

wb.save(XLSX)
print("Overrides applied:")
for t, old, new in applied:
    print(f"  {t}: {old} -> {new}")
print(f"\nMapping_Overrides rows: {len(rows)}")
