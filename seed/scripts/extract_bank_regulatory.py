# -*- coding: utf-8 -*-
"""Bank regulatory sheet writer.
Reads raw/bank_regulatory_source.csv (recapturable source of truth) and writes the
Bank_Regulatory sheet. P/TBV recomputed from raw/us_meta.csv price when TBVPS present.
Contract: the source CSV carries ratios (decimal) in ratio columns and NEVER dollar
GIL/PCL/ACL in a ratio column (see maps/bank_fields.yaml).
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill

from common import RAW, XLSX

SRC = os.path.join(RAW, "bank_regulatory_source.csv")
META = os.path.join(RAW, "us_meta.csv")

BR_COLS = ["Ticker", "Company_Name", "CET1_Ratio", "CET1_Approach", "CET1_Requirement_or_Target",
           "Total_Capital_Ratio", "Leverage_Ratio", "NIM_FY2025", "NIM_Q4_2025", "Efficiency_Ratio",
           "Efficiency_Ratio_Definition", "ROAA", "ROAE", "ROTCE", "PCL_or_PCL_Ratio",
           "NCO_Ratio", "NPL_or_GIL_Ratio", "GIL_Term", "ACL_Loans", "Loan_Growth_YoY",
           "Deposit_Growth_YoY", "Loan_to_Deposit", "TBVPS", "P_TBV", "Dividend_Payout",
           "Share_Count_Change_YoY", "Fiscal_Period_End", "Period_Type", "Source_URL",
           "Page_or_Table", "Confidence", "Notes"]


def write():
    df = pd.read_csv(SRC, keep_default_na=False)
    meta = pd.read_csv(META, index_col=0) if os.path.exists(META) else pd.DataFrame()
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Bank_Regulatory"]
    hfill = PatternFill("solid", fgColor="1F4E78"); hfont = Font(bold=True, color="FFFFFF")
    for j, h in enumerate(BR_COLS, start=1):
        c = ws.cell(row=1, column=j, value=h); c.fill = hfill; c.font = hfont
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for _, row in df.iterrows():
        t = row["Ticker"]
        tbvps = row.get("TBVPS")
        p_tbv = None
        if t in meta.index and pd.notna(meta.loc[t, "price"]) and pd.notna(tbvps) and tbvps not in (None, ""):
            try:
                p_tbv = round(float(meta.loc[t, "price"]) / float(tbvps), 2)
            except (ValueError, ZeroDivisionError):
                p_tbv = None
        for j, h in enumerate(BR_COLS, start=1):
            v = row.get(h, "")
            if h == "P_TBV":
                v = p_tbv
            ws.cell(row=r, column=j, value=v if v != "" else None)
        r += 1
    wb.save(XLSX)
    return len(df)


if __name__ == "__main__":
    print("Bank_Regulatory rows written:", write())
