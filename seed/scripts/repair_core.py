# -*- coding: utf-8 -*-
"""One-shot repair v2: recover missing tickers + rebuild Core_Financials/Time_Series.
- Re-extracts US missing (AEP/XOM/HONA + BRK-B/BF-B hyphen) and CA (ACO.X.TO/IIP-UN.TO + 9 dual-class).
- Disambiguates 7 US/CA bare-ticker collisions (CA keeps .TO).
- Derives Total_Debt (LTD+STD) and Book_Equity (TotalEquity) fallbacks.
- Adds Extraction_Status. Preserves derived columns by Company_Name join. Data only.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) + "/scripts")
sys.path.insert(0, "scripts")
import pandas as pd
import openpyxl

from common import XLSX, save_workbook_atomic
import extract_edgar, extract_ca_ir, assemble

# 1. Re-extract US missing (SEC ticker format: hyphen for dual-class)
US_MISSING = ['AEP', 'XOM', 'HONA', 'BRK-B', 'BF-B']
try:
    adf, tdf = extract_edgar.extract(US_MISSING)
    extract_edgar.upsert(adf, tdf)
    print(f"[repair] US re-extract: {len(adf)} annual rows")
except Exception as e:
    print(f"[repair] US re-extract error: {type(e).__name__}: {e}")

# 2. Re-extract CA missing + dual-class (dot->hyphen fix in extract_ca_ir)
CA_MISSING = ['ACO.X.TO', 'IIP-UN.TO', 'RCI.B.TO', 'QBR.B.TO', 'EMP.A.TO', 'CTC.A.TO',
              'CCL.B.TO', 'TCL.A.TO', 'GIB.A.TO', 'TECK.B.TO', 'BBD.B.TO']
try:
    df, mdf = extract_ca_ir.extract(CA_MISSING)
    extract_ca_ir.upsert(df, mdf)
    print(f"[repair] CA re-extract: {len(df)} rows, {len(mdf)} meta")
except Exception as e:
    print(f"[repair] CA re-extract error: {type(e).__name__}: {e}")

# 3. Build fresh (disambiguated + fallback derivations + Extraction_Status)
core = assemble.build_core()
ts = assemble.build_time_series()
print(f"[repair] core rows={len(core)} unique={core['Ticker'].nunique()} | ts rows={len(ts)}")

# 4. Preserve derived columns from existing (by Company_Name; Rogers special-cased)
existing = pd.read_excel(XLSX, sheet_name='Core_Financials', keep_default_na=False)
existing_cols = list(existing.columns)
derived_cols = [c for c in existing_cols if c not in assemble.CORE_COLS]
name_map = {}
for _, r in existing.iterrows():
    cn = str(r['Company_Name'])
    key = 'Rogers Communications Inc.' if cn == 'Rogers Communications' else cn
    cur = name_map.setdefault(key, {})
    cand = {c: r[c] for c in derived_cols if pd.notna(r[c])}
    if len(cand) > len(cur):
        name_map[key] = cand

# 5. Full-replace Core_Financials
wb = openpyxl.load_workbook(XLSX)
ws = wb['Core_Financials']
hdr = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(hdr)}
if 'Extraction_Status' not in col:
    ws.cell(row=1, column=len(hdr) + 1, value='Extraction_Status')
    col['Extraction_Status'] = len(hdr) + 1
    hdr = hdr + ['Extraction_Status']
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

def clean(v):
    return None if pd.isna(v) else v

written = 0
for _, row in core.iterrows():
    extra = name_map.get(str(row['Company_Name']), {})
    vals = []
    for h in hdr:
        if h == 'Extraction_Status':
            vals.append(row['Extraction_Status'])
        elif h in assemble.CORE_COLS and h in row.index:
            vals.append(clean(row[h]))
        elif h in extra:
            vals.append(clean(extra[h]))
        else:
            vals.append(None)
    ws.append(vals)
    written += 1
print(f"[repair] Core_Financials rows written: {written}")

# 6. Full-replace Time_Series (disambiguated, bulk)
ts_ws = wb['Time_Series']
if ts_ws.max_row > 1:
    ts_ws.delete_rows(2, ts_ws.max_row - 1)
for row in ts.itertuples(index=False, name=None):
    ts_ws.append([clean(v) for v in row])

save_workbook_atomic(wb)
print("[repair] saved atomically")
